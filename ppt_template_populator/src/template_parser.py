from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4
from pptx import Presentation
from .security import validate_pptx
from .target_metadata import ROLE_MIN_CAPACITY, NON_REPLACEABLE_ROLES, normalize_template_targets

EMU_PER_INCH = 914400
EXCLUDED_PLACEHOLDER_TYPES = {"DATE", "FOOTER", "SLIDE_NUMBER"}
DECORATIVE_NAME_HINTS = ("logo", "icon", "decoration", "decorative", "background", "ornament")
INSTRUCTION_TEXT_RE = re.compile(
    r"^\s*visual\s*$|(?:^|\b)(?:add|insert|replace(?:\s+with\s+relevant\s+visual)?|describe|explain|summari[sz]e|set\s+the\s+context|"
    r"clearly\s+define|type|enter|write)\b|\b\d+\s*[-\u2013]\s*\d+\s+(?:points?|bullets?)\b|"
    r"^\s*[\[<{(].+(?:\]|>|}|\))\s*$", re.IGNORECASE,
)
SAMPLE_TEXT_RE = re.compile(
    r"\blorem\s+ipsum\b|\bsample\s+(?:text|content|heading|title)\b|\bplaceholder\b",
    re.IGNORECASE,
)


def classify_existing_text(text: str, role: str, *, is_placeholder: bool = False) -> str:
    normalized = " ".join(str(text).split())
    if role in {"footer", "slide_number"}: return role
    if INSTRUCTION_TEXT_RE.search(normalized): return "instructional_placeholder"
    if SAMPLE_TEXT_RE.search(normalized): return "editable_sample"
    if role in {"decorative", "static", "non_writable", "short_label", "metric", "caption"}:
        return "decorative_or_static"
    if role == "title": return "semantic_heading"
    if is_placeholder: return "replaceable_placeholder"
    return "sample_content" if normalized else "replaceable_placeholder"
STATIC_TEXT_RE = re.compile(r"(?:https?://|www\.|[\w.+-]+@[\w.-]+\.[a-z]{2,}|(?:tel|phone|mobile)\s*[:+]|©|copyright)", re.IGNORECASE)


def _placeholder_type(shape: Any) -> str | None:
    if not getattr(shape, "is_placeholder", False): return None
    try:
        value = shape.placeholder_format.type
        return getattr(value, "name", str(value))
    except (AttributeError, ValueError): return "UNKNOWN"


def _font_size(shape: Any) -> float | None:
    if not getattr(shape, "has_text_frame", False): return None
    sizes = []
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.font.size is not None:
            sizes.append(float(paragraph.font.size.pt))
        for run in paragraph.runs:
            if run.font.size is not None: sizes.append(float(run.font.size.pt))
    return max(sizes) if sizes else None


def _font_name(shape: Any) -> str | None:
    if not getattr(shape, "has_text_frame", False): return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name: return str(run.font.name)
    return None

def _has_bullets(shape: Any) -> bool:
    if not getattr(shape, "has_text_frame", False): return False
    return any(paragraph._p.pPr is not None and any(child.tag.rsplit("}", 1)[-1] in {"buChar", "buAutoNum"} for child in paragraph._p.pPr) for paragraph in shape.text_frame.paragraphs)


def _max_content_length(shape: Any, role: str = "unknown", font_size: float | None = None, existing_text: str = "") -> int:
    width = max(float(shape.width) / EMU_PER_INCH, .25)
    height = max(float(shape.height) / EMU_PER_INCH, .15)
    margins = getattr(shape, "margins", {})
    width = max(.2, width - (margins.get("left", 0) + margins.get("right", 0)) / EMU_PER_INCH)
    height = max(.12, height - (margins.get("top", 0) + margins.get("bottom", 0)) / EMU_PER_INCH)
    points = font_size or 18
    chars_per_line = max(4, width * 72 / max(points * .52, 1))
    lines = max(1, int(height * 72 / max(points * 1.2, 1)))
    geometric = min(3000, int(chars_per_line * lines * .9))
    minimum = ROLE_MIN_CAPACITY.get(role, 20) if lines >= 2 or role in {"title", "subtitle"} else 1
    return max(ROLE_MIN_CAPACITY.get(role, minimum), geometric)


def _capacity(shape: dict[str, Any]) -> dict[str, Any]:
    width = max(.05, shape["position"]["width"] / EMU_PER_INCH)
    height = max(.05, shape["position"]["height"] / EMU_PER_INCH)
    margins = shape.get("text_box_margins") or {}
    available_width = max(.05, width - (margins.get("left", 0) + margins.get("right", 0)) / EMU_PER_INCH)
    available_height = max(.05, height - (margins.get("top", 0) + margins.get("bottom", 0)) / EMU_PER_INCH)
    font_size = float(shape.get("font_size") or 18.0)
    metrics = shape.get("text_metrics") or {}
    line_spacing = float(metrics.get("line_spacing") or 1.2)
    paragraph_spacing = float(metrics.get("paragraph_spacing_pt") or 0)
    bullet_indent = float(metrics.get("bullet_indent_inches") or 0)
    available_width = max(.05, available_width - bullet_indent)
    line_height = max(font_size * line_spacing + paragraph_spacing, 1)
    maximum_lines = max(1, int(available_height * 72 / line_height))
    characters_per_line = max(4, int(available_width * 72 / max(font_size * .52, 1)))
    maximum_characters = max(1, min(3000, int(characters_per_line * maximum_lines * .9)))
    maximum_words = max(1, int(maximum_characters / 6.2))
    existing = str(shape.get("existing_text") or "")
    estimated_existing_lines = sum(max(1, (len(line) + characters_per_line - 1) // characters_per_line) for line in (existing.splitlines() or [""]))
    return {"available_width": available_width, "available_height": available_height,
            "characters_per_line": characters_per_line, "estimated_line_count": estimated_existing_lines,
            "maximum_lines": maximum_lines, "maximum_words": maximum_words,
            "maximum_characters": maximum_characters, "line_spacing": line_spacing,
            "paragraph_spacing_pt": paragraph_spacing, "bullet_indent_inches": bullet_indent,
            "font_size": font_size, "replacement_state": True}


def _margins(shape: Any) -> dict[str, int]:
    frame = getattr(shape, "text_frame", None)
    return {name: int(getattr(frame, f"margin_{name}", 0) or 0) for name in ("left", "right", "top", "bottom")} if frame else {}


def _text_metrics(shape: Any) -> dict[str, float]:
    if not getattr(shape, "has_text_frame", False):
        return {}
    paragraphs = list(shape.text_frame.paragraphs)
    spacings = [value for paragraph in paragraphs for value in (
        getattr(paragraph, "space_before", None), getattr(paragraph, "space_after", None)
    ) if value is not None]
    line_values = [paragraph.line_spacing for paragraph in paragraphs
                   if isinstance(paragraph.line_spacing, (int, float))]
    maximum_level = max((int(paragraph.level or 0) for paragraph in paragraphs), default=0)
    return {
        "line_spacing": float(max(line_values, default=1.2)),
        "paragraph_spacing_pt": float(max((value.pt for value in spacings), default=0.0)),
        "bullet_indent_inches": .25 * maximum_level,
    }


def _walk_shapes(shapes: Iterable[Any], group_path: tuple[int, ...] = ()) -> Iterable[tuple[Any, tuple[int, ...]]]:
    for shape in shapes:
        path = (*group_path, int(shape.shape_id))
        yield shape, path
        children = getattr(shape, "shapes", None)
        if children is not None: yield from _walk_shapes(children, path)


def _explicitly_replaceable(shape: dict[str, Any]) -> bool:
    name = shape["shape_name"].lower()
    return "replaceable" in name or name.startswith("content_") or name.startswith("content ")


def _classify(shape: dict[str, Any], slide_width: int, slide_height: int, largest_font: float) -> str:
    text = shape["existing_text"].strip(); placeholder_type = (shape["placeholder_type"] or "").upper()
    if placeholder_type in {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}: return "title"
    if placeholder_type in {"SUBTITLE"}: return "subtitle"
    if placeholder_type in {"BODY", "OBJECT", "VERTICAL_BODY"}: return "bullets" if shape.get("has_bullets") else "body"
    if placeholder_type == "SLIDE_NUMBER": return "slide_number"
    if placeholder_type in {"DATE", "FOOTER"}: return "footer"
    name = shape["shape_name"].lower()
    if any(hint in name for hint in DECORATIVE_NAME_HINTS): return "decorative"
    if STATIC_TEXT_RE.search(text): return "non_writable"
    if re.fullmatch(r"\s*(?:\d+|\d+\s*/\s*\d+|page\s+\d+)\s*", text, re.IGNORECASE):
        tiny_label = shape["position"]["width"] <= slide_width * .08 and shape["position"]["height"] <= slide_height * .08
        return "slide_number" if shape["position"]["top"] >= slide_height * .75 or tiny_label else "metric"
    top = shape["position"]["top"]; bottom = top + shape["position"]["height"]
    font = shape["font_size"] or 0
    if bottom >= slide_height * .92 and font <= 14: return "footer"
    if font >= max(28, largest_font * .9) and top <= slide_height * .45: return "title"
    if font >= 22 and len(text) <= 120: return "title"
    position = shape["position"]
    if (position["top"] <= slide_height * .24 and
            EMU_PER_INCH * .55 <= position["height"] <= EMU_PER_INCH * 1.4 and
            position["width"] >= slide_width * .4 and len(text) <= 120):
        return "title"
    if shape["position"]["height"] <= int(EMU_PER_INCH * .55): return "short_label"
    if shape["position"]["height"] >= slide_height * .12:
        return "bullets" if shape.get("has_bullets") or text.count("\n") >= 1 else "body"
    if font >= 16 and top <= slide_height * .55 and len(text) <= 160: return "subtitle"
    if font and font <= 12 and (shape["position"]["width"] <= slide_width * .55 or len(text) <= 100): return "caption"
    if shape.get("has_bullets"): return "bullets"
    if len(text) >= 20 or shape["position"]["height"] >= slide_height * .12: return "body"
    if text: return "static"
    return "non_writable"


def _writable(shape: dict[str, Any]) -> bool:
    if not shape["has_text_frame"]: return False
    explicit = _explicitly_replaceable(shape)
    if shape["role"] in NON_REPLACEABLE_ROLES and not explicit: return False
    if shape["is_placeholder"]:
        return shape["placeholder_idx"] is not None and ((shape["placeholder_type"] or "").upper() not in EXCLUDED_PLACEHOLDER_TYPES or explicit)
    text = shape["existing_text"].strip()
    return bool(text and re.search(r"[A-Za-z0-9]", text))


def _classification_confidence(shape: dict[str, Any]) -> float:
    if shape["is_placeholder"] or shape["role"] in {"title", "body", "bullets", "footer", "slide_number", "decorative", "non_writable"}: return .95
    if shape["role"] in {"subtitle", "caption", "label", "metric"}: return .75
    return .35

SECTION_PATTERNS = (
    ("thank_you", ("thank you", "questions", "q&a", "discussion")),
    ("introduction", ("introduction", "overview", "context", "about")),
    ("agenda", ("agenda", "contents", "outline")),
    ("problem_statement", ("problem", "challenge", "pain point")),
    ("objectives", ("objective", "goal", "aim")),
    ("proposed_solution", ("solution", "framework", "proposal")),
    ("methodology_workflow", ("method", "workflow", "process", "approach", "how it works")),
    ("findings_evidence", ("finding", "result", "evidence", "analysis", "metric")),
    ("conclusion", ("conclusion", "summary", "takeaway")),
    ("case_study", ("case study", "example")),
    ("roadmap", ("roadmap", "timeline", "next step")),
    ("risks", ("risk", "limitation", "constraint")),
    ("team", ("team", "people", "leadership")),
)

def _detect_section(slide: dict[str, Any], number: int, slide_count: int) -> str:
    title_text = " ".join(target.get("existing_text", "") for target in slide.get("targets", []) if target.get("role") == "title")
    evidence = f"{title_text} {slide.get('layout_name', '')}".casefold()
    for section, patterns in SECTION_PATTERNS:
        if any(pattern in evidence for pattern in patterns): return section
    if number == 1: return "title"
    if number == slide_count: return "thank_you"
    return "content"


def parse_template(pptx_path: Path, *, file_path: str, template_name: str | None = None, description: str = "", category: str = "General", template_id: str | None = None, max_bytes: int = 25 * 1024 * 1024) -> dict[str, Any]:
    validate_pptx(pptx_path, max_bytes); presentation = Presentation(str(pptx_path)); slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        shapes = []
        for shape, group_path in _walk_shapes(slide.shapes):
            placeholder_idx = None
            if shape.is_placeholder:
                try: placeholder_idx = int(shape.placeholder_format.idx)
                except (AttributeError, ValueError): pass
            has_text = bool(getattr(shape, "has_text_frame", False)); font_size = _font_size(shape)
            shapes.append({"slide_number": number, "shape_id": int(shape.shape_id), "shape_name": shape.name, "group_path": list(group_path[:-1]), "is_placeholder": bool(shape.is_placeholder), "placeholder_idx": placeholder_idx, "placeholder_type": _placeholder_type(shape), "existing_text": shape.text if has_text else "", "has_text_frame": has_text, "has_bullets": _has_bullets(shape), "font_size": font_size, "font_name": _font_name(shape), "max_content_length": 0, "text_box_margins": _margins(shape), "text_metrics": _text_metrics(shape), "position": {"left": int(shape.left), "top": int(shape.top), "width": int(shape.width), "height": int(shape.height)}})
        largest_font = max((item["font_size"] or 0 for item in shapes), default=0)
        targets = []
        for item in shapes:
            item["role"] = _classify(item, int(presentation.slide_width), int(presentation.slide_height), largest_font)
            item["text_category"] = classify_existing_text(
                item["existing_text"], item["role"], is_placeholder=item["is_placeholder"]
            )
            item["classification_confidence"] = _classification_confidence(item)
            item["max_content_length"] = _max_content_length(SimpleShape(item), item["role"], item["font_size"], item["existing_text"]) if item["has_text_frame"] else 0
            item["capacity"] = _capacity(item) if item["has_text_frame"] else {}
            if item["has_text_frame"]:
                item["max_content_length"] = item["capacity"]["maximum_characters"]
            item["is_writable_target"] = _writable(item) or (
                item["text_category"] in {"instructional_placeholder", "editable_sample", "sample_content"}
                and item["role"] not in NON_REPLACEABLE_ROLES
            )
            meaningful_protected_text = bool(item["existing_text"].strip() and re.search(r"[A-Za-z0-9]", item["existing_text"]))
            if item["is_writable_target"] or meaningful_protected_text:
                targets.append({"slide_number": number, "target_kind": "placeholder" if item["is_placeholder"] else "shape", "target_id": item["placeholder_idx"] if item["is_placeholder"] else item["shape_id"], "role": item["role"], "existing_text": item["existing_text"], "text_category": item["text_category"], "replacement_mode": "replace", "clear_before_write": True, "approximate_max_characters": item["max_content_length"], "max_content_length": item["max_content_length"], "capacity": item["capacity"], "text_box_margins": item["text_box_margins"], "text_metrics": item["text_metrics"], "required": bool(item["is_placeholder"] and item["is_writable_target"]), "replaceable": bool(item["is_writable_target"]), "classification_confidence": item["classification_confidence"], "shape_name": item["shape_name"], "font_size": item["font_size"], "font_name": item["font_name"], "position": item["position"]})
        placeholder_types = {(item["placeholder_type"] or "").upper() for item in shapes}
        slides.append({"slide_number": number, "layout_name": slide.slide_layout.name or "Unnamed layout", "has_image_placeholders": bool(placeholder_types & {"BITMAP", "PICTURE", "MEDIA_CLIP"}), "has_chart_placeholders": bool(placeholder_types & {"CHART", "OBJECT"}), "shapes": shapes, "targets": targets})
    occurrences = {}
    for slide in slides:
        for shape in slide["shapes"]:
            text = " ".join(shape["existing_text"].casefold().split())
            if text and len(text) <= 80: occurrences.setdefault(text, set()).add(slide["slide_number"])
    for slide in slides:
        for shape in slide["shapes"]:
            text = " ".join(shape["existing_text"].casefold().split())
            if text and len(occurrences.get(text, ())) > 1 and shape["role"] not in {"title", "subtitle", "body", "bullets", "slide_number", "footer"}:
                shape["role"], shape["is_writable_target"] = "static", False
        slide["section_type"] = _detect_section(slide, slide["slide_number"], len(slides))
        if slide["section_type"] in {"title", "thank_you"} and not any(shape["role"] == "title" and shape["is_writable_target"] for shape in slide["shapes"]):
            candidates = [shape for shape in slide["shapes"] if shape["is_writable_target"] and shape["existing_text"].strip()]
            if candidates:
                best = max(candidates, key=lambda shape: (
                    shape["position"]["width"] * shape["position"]["height"],
                    -shape["position"]["top"],
                ))
                best["role"] = "title"
        title_shapes = sorted((shape for shape in slide["shapes"] if shape["role"] == "title" and shape["is_writable_target"]), key=lambda shape: (-(shape["font_size"] or 0), shape["position"]["top"]))
        primary_title = title_shapes[0] if title_shapes else None
        for shape in slide["shapes"]:
            if shape in title_shapes[1:]: shape["role"], shape["is_writable_target"] = "static", False
            if shape["role"] == "subtitle" and slide["section_type"] not in {"title", "thank_you"}:
                shape["role"], shape["is_writable_target"] = "static", False
            if slide["section_type"] == "agenda" and shape is not primary_title:
                shape["role"], shape["is_writable_target"] = "static", False
        shape_by_key = {("placeholder" if shape["is_placeholder"] else "shape", shape["placeholder_idx"] if shape["is_placeholder"] else shape["shape_id"]): shape for shape in slide["shapes"]}
        slide["targets"] = [
            target for target in slide["targets"]
            if shape_by_key[(target["target_kind"], target["target_id"])]["is_writable_target"]
            or target.get("text_category") in {"instructional_placeholder", "editable_sample", "sample_content"}
            or len(occurrences.get(" ".join(target["existing_text"].casefold().split()), ())) > 1
        ]
        for target in slide["targets"]:
            shape = shape_by_key[(target["target_kind"], target["target_id"])]
            target["role"] = shape["role"]
            target["replaceable"] = bool(shape["is_writable_target"])
            target["required"] = bool(target["required"] and shape["is_writable_target"])
        title = next((target for target in slide["targets"] if target["replaceable"] and target["role"] == "title"), None)
        bodies = [target for target in slide["targets"] if target["replaceable"] and target["role"] in {"body", "bullets"}]
        capacity_summary = {f"{target['target_kind']}:{target['target_id']}": target["approximate_max_characters"] for target in slide["targets"] if target["replaceable"]}
        slide.update({"title_target": ({"target_kind": title["target_kind"], "target_id": title["target_id"]} if title else None), "body_targets": [{"target_kind": target["target_kind"], "target_id": target["target_id"]} for target in bodies], "visual_targets": [{"shape_id": shape["shape_id"], "role": shape["role"]} for shape in slide["shapes"] if shape["role"] in {"metric", "caption", "short_label", "decorative", "static"}], "writable_targets": [{"target_kind": target["target_kind"], "target_id": target["target_id"]} for target in slide["targets"] if target["replaceable"]], "static_targets": [{"shape_id": shape["shape_id"], "role": shape["role"]} for shape in slide["shapes"] if not shape["is_writable_target"] and shape["has_text_frame"]], "capacity_summary": capacity_summary, "text_capacity": capacity_summary})
    usable_placeholders = sum(target["target_kind"] == "placeholder" and target["replaceable"] for slide in slides for target in slide["targets"])
    usable_shapes = sum(target["target_kind"] == "shape" and target["replaceable"] for slide in slides for target in slide["targets"])
    slides_without = sum(not any(target["replaceable"] for target in slide["targets"]) for slide in slides); now = datetime.now(timezone.utc).isoformat()
    return normalize_template_targets({"template_id": template_id or uuid4().hex, "template_name": (template_name or pptx_path.stem).strip(), "description": description.strip(), "category": category.strip() or "General", "file_path": file_path, "slide_width": int(presentation.slide_width), "slide_height": int(presentation.slide_height), "slide_count": len(slides), "slides": slides, "has_image_placeholders": any(slide["has_image_placeholders"] for slide in slides), "has_chart_placeholders": any(slide["has_chart_placeholders"] for slide in slides), "usable_placeholder_count": usable_placeholders, "usable_shape_target_count": usable_shapes, "slides_without_writable_targets": slides_without, "safe_for_automatic_population": bool(usable_placeholders + usable_shapes) and slides_without == 0, "created_at": now, "updated_at": now})


class SimpleShape:
    def __init__(self, metadata: dict[str, Any]):
        self.width = metadata["position"]["width"]
        self.height = metadata["position"]["height"]
        self.margins = metadata.get("text_box_margins", {})
