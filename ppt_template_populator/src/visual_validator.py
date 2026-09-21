from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from .target_metadata import replaceable_targets_by_slide, required_targets_by_slide
from .template_binary import duplicate_package_members
from .ppt_populator import _contrast, _effective_background, _estimated_lines, _resolved_rgb, _theme_colors


@dataclass(frozen=True)
class VisualIssue:
    slide_number: int
    category: str
    message: str


def validate_presentation_layout(path: Path, template: dict[str, Any]) -> list[VisualIssue]:
    """Perform renderer-independent structural visual checks on the saved PPTX."""
    presentation = Presentation(str(path))
    issues: list[VisualIssue] = []
    duplicates = duplicate_package_members(path)
    if duplicates:
        issues.append(VisualIssue(0, "duplicate_package_member", f"The PPTX contains {len(duplicates)} duplicate package member name(s)."))
    theme = _theme_colors(presentation)
    expected_count = int(template.get("slide_count") or len(template.get("slides", [])))
    if len(presentation.slides) != expected_count:
        issues.append(VisualIssue(0, "slide_count", f"Expected {expected_count} slides but the output contains {len(presentation.slides)}."))
    required = required_targets_by_slide(template)
    replaceable = replaceable_targets_by_slide(template)
    planned_titles = {
        slide["slide_number"]: str(slide.get("planned_section_title") or "").strip().casefold()
        for slide in template.get("slides", [])
    }
    expected_samples = {
        slide["slide_number"]: {str(target.get("existing_text", "")).strip() for target in slide.get("targets", []) if target.get("replaceable") and str(target.get("existing_text", "")).strip()}
        for slide in template.get("slides", [])
    }
    for number, slide in enumerate(presentation.slides, start=1):
        def walk(shapes):
            for shape in shapes:
                yield shape
                children = getattr(shape, "shapes", None)
                if children is not None: yield from walk(children)
        text_shapes = [shape for shape in walk(slide.shapes)
                       if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        visible_text = "\n".join(shape.text for shape in text_shapes)
        placeholder_map, shape_map = {}, {}
        for shape in walk(slide.shapes):
            if shape.is_placeholder:
                try: placeholder_map[int(shape.placeholder_format.idx)] = shape
                except (AttributeError, ValueError): pass
            else: shape_map[int(shape.shape_id)] = shape
        editable_shape_ids = {
            shape.shape_id
            for kind, target_id in replaceable.get(number, {})
            for shape in [placeholder_map.get(target_id) if kind == "placeholder" else shape_map.get(target_id)]
            if shape is not None
        }
        for (kind, target_id), metadata in required.get(number, {}).items():
            shape = placeholder_map.get(target_id) if kind == "placeholder" else shape_map.get(target_id)
            if shape is None or not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                issues.append(VisualIssue(number, "empty_required_target", f"Required {metadata.get('role', 'text')} target is empty."))
        for (kind, target_id), metadata in replaceable.get(number, {}).items():
            shape = placeholder_map.get(target_id) if kind == "placeholder" else shape_map.get(target_id)
            if shape is None or not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            original = " ".join(str(metadata.get("existing_text") or "").split()).casefold()
            current = " ".join(shape.text.split()).casefold()
            if metadata.get("text_category") in {"instructional_placeholder", "editable_sample", "sample_content"} and current == original:
                issues.append(VisualIssue(
                    number, "unresolved_instruction",
                    f"{kind.title()} {target_id} ({metadata.get('role', 'text')}) still contains {metadata.get('text_category')}."
                ))
            if metadata.get("role") == "title":
                heading = " ".join(shape.text.split()).strip()
                if not heading or len(heading) == 1 or heading.endswith(("-", "/", ":")):
                    issues.append(VisualIssue(number, "incomplete_heading", f"{kind.title()} {target_id} contains an incomplete heading."))
            capacity = metadata.get("capacity") or {}
            if capacity and _estimated_lines(shape.text, capacity) > int(capacity.get("maximum_lines") or 1):
                issues.append(VisualIssue(number, "overflow", f"Editable {metadata.get('role', 'text')} target exceeds its vertical line capacity."))
            background = _effective_background(shape, slide, theme)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    foreground = _resolved_rgb(run.font.color, theme)
                    if foreground is None: continue
                    size = run.font.size.pt if run.font.size else float(metadata.get("font_size") or 18)
                    required_ratio = 3.0 if size >= 18 or bool(run.font.bold and size >= 14) else 4.5
                    if _contrast(foreground, background) < required_ratio:
                        issues.append(VisualIssue(number, "contrast", f"Editable text does not meet the {required_ratio:g}:1 contrast target."))
                        break
        template_slide = next((item for item in template.get("slides", [])
                               if item.get("slide_number") == number), {})
        for metadata in template_slide.get("targets", []):
            if not metadata.get("clear_before_write") or metadata.get("text_category") not in {
                "replaceable_placeholder", "instructional_placeholder", "editable_sample", "sample_content"
            }:
                continue
            kind, target_id = metadata.get("target_kind"), metadata.get("target_id")
            shape = placeholder_map.get(target_id) if kind == "placeholder" else shape_map.get(target_id)
            if shape is None or not getattr(shape, "has_text_frame", False):
                continue
            original = " ".join(str(metadata.get("existing_text") or "").split()).casefold()
            current = " ".join(shape.text.split()).casefold()
            if current and current == original:
                issues.append(VisualIssue(
                    number, "unresolved_instruction",
                    f"{str(kind).title()} {target_id} ({metadata.get('role', 'text')}) still contains {metadata.get('text_category')}."
                ))
        if "pptxgenjs presentation" in visible_text.casefold():
            issues.append(VisualIssue(number, "forbidden_title", "Generic PptxGenJS title remains in the output."))
        assignment = template_slide.get("slide_assignment") or {}
        if assignment.get("section_type") == "introduction":
            title_ref = assignment.get("title_target") or {}
            body_ref = assignment.get("primary_body_target") or {}
            def assigned_shape(reference):
                target_id = reference.get("target_id")
                return (placeholder_map.get(target_id) if reference.get("target_kind") == "placeholder"
                        else shape_map.get(target_id))
            title_shape = assigned_shape(title_ref)
            body_shape = assigned_shape(body_ref)
            if title_shape is None or not getattr(title_shape, "has_text_frame", False) or not title_shape.text.strip():
                issues.append(VisualIssue(number, "missing_introduction_heading", "The mandatory Introduction heading is missing from its assigned title target."))
            if body_ref and (body_shape is None or not getattr(body_shape, "has_text_frame", False) or not body_shape.text.strip()):
                issues.append(VisualIssue(number, "missing_introduction_body", "The mandatory Introduction body is missing from its assigned primary body target."))
            if str(template_slide.get("section_type")) != "introduction":
                issues.append(VisualIssue(number, "incorrect_semantic_mapping", "The Introduction assignment does not match the planned slide semantics."))
        if number == len(presentation.slides) and "thank you" not in visible_text.casefold():
            issues.append(VisualIssue(number, "section_order", "The final slide is not the mandatory Thank You closing."))
        for shape in text_shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height:
                issues.append(VisualIssue(number, "bounds", f"Text shape {shape.shape_id} extends outside the slide."))
            normalized_text = shape.text.strip().casefold()
            if shape.shape_id in editable_shape_ids and shape.text.strip() in expected_samples.get(number, set()) and normalized_text != planned_titles.get(number):
                issues.append(VisualIssue(number, "sample_text", f"Editable sample text remains in shape {shape.shape_id}."))
            sizes = [run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size]
            if shape.shape_id in editable_shape_ids and sizes and min(sizes) < 14:
                issues.append(VisualIssue(number, "font_size", f"Text shape {shape.shape_id} uses text below 14 pt."))
        for index, left in enumerate(text_shapes):
            for right in text_shapes[index + 1:]:
                overlap_width = min(left.left + left.width, right.left + right.width) - max(left.left, right.left)
                overlap_height = min(left.top + left.height, right.top + right.height) - max(left.top, right.top)
                if overlap_width > 0 and overlap_height > 0:
                    overlap = overlap_width * overlap_height
                    smaller = max(1, min(left.width * left.height, right.width * right.height))
                    if overlap / smaller > .2:
                        issues.append(VisualIssue(number, "collision", f"Text shapes {left.shape_id} and {right.shape_id} materially overlap."))
    return issues
