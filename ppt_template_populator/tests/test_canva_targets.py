from pathlib import Path
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from src.ppt_populator import populate_presentation
from src.response_models import PresentationContent
from src.response_validator import ResponseValidationError, validate_response
from src.template_parser import parse_template


def _textbox(collection, left, top, width, height, text, size=18, name=None):
    shape = collection.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name: shape.name = name
    run = shape.text_frame.paragraphs[0].add_run(); run.text = text; run.font.size = Pt(size); run.font.bold = True; run.font.color.rgb = RGBColor(12, 34, 56)
    return shape


@pytest.fixture
def canva_pptx(tmp_path: Path) -> Path:
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide.shapes, 1, .5, 8, 1, "Startup Business", 32)
    _textbox(slide.shapes, 1, 2, 8, 3, "A meaningful body paragraph for an investor presentation.", 16)
    _textbox(slide.shapes, .2, 7.1, 2, .2, "Confidential", 9)
    _textbox(slide.shapes, 9.2, 7.1, .3, .2, "1", 9)
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(.2), Inches(.2))
    group = slide.shapes.add_group_shape()
    _textbox(group.shapes, 1, 5.2, 5, .6, "Grouped Canva callout text", 20)
    path = tmp_path / "canva.pptx"; prs.save(path); return path


def test_standard_placeholders_remain_placeholder_targets(template_metadata):
    targets = template_metadata["slides"][0]["targets"]
    assert targets and all(target["target_kind"] == "placeholder" for target in targets)
    assert {target["target_id"] for target in targets} == {0, 1}


def test_canva_textboxes_and_grouped_shapes_are_targets(canva_pptx):
    metadata = parse_template(canva_pptx, file_path="canva.pptx")
    targets = metadata["slides"][0]["targets"]
    texts = {target["existing_text"]: target for target in targets}
    assert texts["Startup Business"]["target_kind"] == "shape"
    assert texts["Startup Business"]["role"] == "title"
    assert "Grouped Canva callout text" in texts
    grouped_shape = next(shape for shape in metadata["slides"][0]["shapes"] if shape["existing_text"] == "Grouped Canva callout text")
    assert grouped_shape["group_path"]


def test_empty_decorative_footer_and_page_number_are_not_targets(canva_pptx):
    metadata = parse_template(canva_pptx, file_path="canva.pptx")
    texts = {target["existing_text"] for target in metadata["slides"][0]["targets"]}
    assert "" not in texts and "Confidential" not in texts and "1" not in texts


def test_invented_canva_shape_id_is_rejected(canva_pptx):
    metadata = parse_template(canva_pptx, file_path="canva.pptx")
    payload = {"presentation_title": "Test", "slides": [{"slide_number": 1, "layout_name": "Blank", "targets": [{"target_kind": "shape", "target_id": 99999, "content": "Invented"}]}]}
    with pytest.raises(ResponseValidationError, match="Unknown shape"): validate_response(payload, metadata)


def test_canva_population_applies_standard_typography_and_preserves_color(canva_pptx, tmp_path):
    # The template-derived typography profile preserves its font family and
    # clamps the title into the readable reference hierarchy.
    metadata = parse_template(canva_pptx, file_path="canva.pptx")
    target = next(item for item in metadata["slides"][0]["targets"] if item["existing_text"] == "Startup Business")
    assert target["role"] == "title"
    payload = {"presentation_title": "Test", "slides": [{"slide_number": 1, "layout_name": "Blank", "targets": [{"target_kind": item["target_kind"], "target_id": item["target_id"], "content": "EcoBin AI" if item == target else item["existing_text"]} for item in metadata["slides"][0]["targets"]]}]}
    content, _ = validate_response(payload, metadata); output, _ = populate_presentation(canva_pptx, content, tmp_path / "out", metadata)
    prs = Presentation(output)
    def walk(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"): yield from walk(shape.shapes)
    shape = next(item for item in walk(prs.slides[0].shapes) if item.shape_id == target["target_id"])
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.text == "EcoBin AI" and run.font.name == (target.get("font_name") or "Aptos") and 28 <= run.font.size.pt <= 44 and run.font.bold
    assert run.font.color.rgb == RGBColor(12, 34, 56)


def test_template_compatibility_scoring(canva_pptx, tmp_path):
    compatible = parse_template(canva_pptx, file_path="canva.pptx")
    assert compatible["usable_shape_target_count"] >= 3
    assert compatible["slides_without_writable_targets"] == 0 and compatible["safe_for_automatic_population"]
    prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); empty = tmp_path / "empty.pptx"; prs.save(empty)
    incompatible = parse_template(empty, file_path="empty.pptx")
    assert incompatible["slides_without_writable_targets"] == 1
    assert not incompatible["safe_for_automatic_population"]


def test_minor_overflow_reduces_font_within_safe_minimum(canva_pptx, tmp_path):
    metadata = parse_template(canva_pptx, file_path="canva.pptx")
    target = next(item for item in metadata["slides"][0]["targets"] if item["role"] == "title")
    target["approximate_max_characters"] = target["max_content_length"] = 60
    required = [item for item in metadata["slides"][0]["targets"] if item["required"]]
    payload_targets = [{"target_kind": item["target_kind"], "target_id": item["target_id"], "content": "A" * 66 if item == target else item["existing_text"]} for item in required]
    payload = {"presentation_title": "Test", "slides": [{"slide_number": 1, "layout_name": "Blank", "targets": payload_targets}]}
    content, warnings = validate_response(payload, metadata)
    assert any("minor overflow" in warning for warning in warnings)
    output, population_warnings = populate_presentation(canva_pptx, content, tmp_path / "out", metadata)
    prs = Presentation(output)
    shape = next(item for item in prs.slides[0].shapes if item.shape_id == target["target_id"])
    assert 20 <= shape.text_frame.paragraphs[0].runs[0].font.size.pt < 36
    assert any("font size was reduced" in warning for warning in population_warnings)
