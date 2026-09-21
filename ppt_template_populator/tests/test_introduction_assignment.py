from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from src.ppt_populator import populate_presentation
from src.response_models import PresentationContent, SlideContent, TargetContent
from src.semantic_content import validate_semantic_response
from src.slide_planner import build_slide_plan
from src.template_parser import parse_template
from src.visual_validator import validate_presentation_layout


def _source(path: Path) -> tuple[int, int, int]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.7), Inches(.4), Inches(8.5), Inches(.7))
    title.text = "Introduction"; title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    body = slide.shapes.add_textbox(Inches(.8), Inches(1.5), Inches(8), Inches(3.8))
    body.text = "Set the context with 3-5 concise points"
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    unused = slide.shapes.add_textbox(Inches(9.2), Inches(1.6), Inches(3), Inches(1.2))
    unused.text = "Add supporting example"
    brand = slide.shapes.add_textbox(Inches(.5), Inches(7), Inches(2.5), Inches(.2))
    brand.text = "BRAND LEGAL"; brand.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    deck.save(path)
    return title.shape_id, body.shape_id, unused.shape_id


def test_slide_plan_has_explicit_introduction_assignment(tmp_path):
    source = tmp_path / "template.pptx"
    _source(source)
    metadata = parse_template(source, file_path="templates/template.pptx")
    planned = build_slide_plan(metadata, "Grounded source statement. More source context.", 3)
    introduction = planned["slides"][1]
    assignment = introduction["slide_assignment"]
    assert assignment["section_type"] == "introduction"
    assert assignment["template_slide_number"] == introduction["source_slide_number"]
    assert assignment["title_target"] != assignment["primary_body_target"]
    assert assignment["title_target"]["target_id"]
    assert assignment["primary_body_target"]["target_id"]


def test_agenda_is_not_accepted_as_introduction():
    try:
        validate_semantic_response({"slides": [{
            "slide_number": 8, "section_type": "agenda", "title": "Agenda",
            "bullets": ["Grounded overview"], "speaker_notes": None,
        }]}, {8: "introduction"})
    except ValueError as exc:
        assert "does not match 'introduction'" in str(exc)
    else:
        raise AssertionError("Agenda content must not satisfy an Introduction assignment")


def test_population_clears_unused_samples_and_preserves_static_text(tmp_path):
    source = tmp_path / "template.pptx"
    title_id, body_id, unused_id = _source(source)
    metadata = parse_template(source, file_path="templates/template.pptx")
    slide = metadata["slides"][0]
    slide["section_type"] = "introduction"
    slide["planned_section_title"] = "Introduction"
    slide["slide_assignment"] = {
        "template_slide_number": 1, "section_type": "introduction",
        "title_target": {"target_kind": "shape", "target_id": title_id},
        "subtitle_target": None,
        "primary_body_target": {"target_kind": "shape", "target_id": body_id},
        "optional_body_targets": [], "replaceable_targets": [], "static_targets": [],
    }
    generated = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name="Blank", section_type="introduction", targets=[
            TargetContent(target_kind="shape", target_id=title_id, content_type="title", text="Introduction"),
            TargetContent(target_kind="shape", target_id=body_id, content_type="bullets", bullets=["Grounded source context explains the subject"]),
        ],
    )])
    output, _ = populate_presentation(source, generated, tmp_path / "generated", metadata)
    result = Presentation(output)
    shapes = {shape.shape_id: shape for shape in result.slides[0].shapes}
    assert shapes[title_id].text == "Introduction"
    assert "Grounded source context" in shapes[body_id].text
    assert shapes[unused_id].text == ""
    assert any(shape.text == "BRAND LEGAL" for shape in result.slides[0].shapes if shape.has_text_frame)
    categories = {issue.category for issue in validate_presentation_layout(output, metadata)}
    assert not categories & {"missing_introduction_heading", "missing_introduction_body", "unresolved_instruction", "sample_text"}


def test_assignment_validation_uses_metadata_ids_not_fixed_shape_numbers(tmp_path):
    source = tmp_path / "template.pptx"
    title_id, body_id, _ = _source(source)
    metadata = parse_template(source, file_path="templates/template.pptx")
    slide = metadata["slides"][0]
    slide["section_type"] = "introduction"
    slide["slide_assignment"] = {
        "template_slide_number": 1, "section_type": "introduction",
        "title_target": {"target_kind": "shape", "target_id": title_id},
        "primary_body_target": {"target_kind": "shape", "target_id": body_id},
    }
    issues = validate_presentation_layout(source, metadata)
    assert not any(issue.category.startswith("missing_introduction") for issue in issues)
