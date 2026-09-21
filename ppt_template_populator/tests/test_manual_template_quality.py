from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from src.pipeline import GenerationPipeline
from src.ppt_populator import populate_presentation
from src.response_models import PresentationContent, SlideContent, TargetContent
from src.semantic_content import shorten_heading
from src.template_binary import checksum_sha256
from src.template_parser import INSTRUCTION_TEXT_RE, classify_existing_text, parse_template
from src.template_profile import TemplateProfile
from src.visual_validator import validate_presentation_layout


def test_instructional_and_sample_text_are_detected_generically():
    for value in ("Add supporting evidence", "Describe the method", "Set the context in 3-5 points",
                  "[Insert title]", "Lorem ipsum dolor"):
        assert classify_existing_text(value, "body") in {"instructional_placeholder", "editable_sample", "sample_content"}
    assert classify_existing_text("Copyright Example Corp", "footer") == "footer"


def test_heading_shortening_uses_complete_words_and_semantic_fallback():
    assert shorten_heading("A Comprehensive Overview of Operational Planning", 22, "Methodology") == "Operational Planning"
    assert shorten_heading("Unreasonablylongheadingword", 12, "Objectives") == "Objectives"
    with pytest.raises(Exception, match="complete slide heading"):
        shorten_heading("Longword", 2, "Gap")


def test_manual_and_automatic_paths_use_same_complete_template_profile(indexed_template, tmp_path):
    class Retriever:
        def get_full(self, template_id): return deepcopy(indexed_template)
    pipeline = GenerationPipeline(Retriever(), SimpleNamespace(), tmp_path, tmp_path)
    manual = pipeline._full_template(indexed_template["template_id"], selection_mode="manual")
    automatic = pipeline._full_template(indexed_template["template_id"], selection_mode="automatic")
    assert TemplateProfile.model_validate(manual) == TemplateProfile.model_validate(automatic)
    for key in ("slides", "checksum_sha256", "minio_bucket", "minio_object_key", "pptx_binary_bytes"):
        assert key in manual


def test_partial_manual_metadata_is_rejected_and_logged(caplog, tmp_path):
    class Retriever:
        def get_full(self, template_id): return {"template_id": template_id, "template_name": "Partial"}
    pipeline = GenerationPipeline(Retriever(), SimpleNamespace(), tmp_path, tmp_path)
    with pytest.raises(Exception):
        pipeline._full_template("partial", selection_mode="manual")
    assert "stage=full_template_retrieval" in caplog.text
    assert "selection_mode=manual" in caplog.text


def test_unused_instruction_is_cleared_and_static_brand_is_preserved(tmp_path):
    source = tmp_path / "instruction-template.pptx"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(8), Inches(.8))
    title.text = "Sample title"; title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    body = slide.shapes.add_textbox(Inches(.8), Inches(1.7), Inches(8), Inches(3.5))
    body.text = "Describe the supporting evidence in 3-5 bullets"
    unused = slide.shapes.add_textbox(Inches(.8), Inches(5.4), Inches(4), Inches(.6))
    unused.text = "Insert an optional explanation"
    brand = slide.shapes.add_textbox(Inches(8), Inches(6.8), Inches(1.5), Inches(.3))
    brand.text = "Example Corp"
    deck.save(source)
    metadata = parse_template(source, file_path="templates/instruction-template.pptx")
    slide_meta = metadata["slides"][0]
    body_target = next(target for target in slide_meta["targets"] if target["target_id"] == body.shape_id)
    title_target = next(target for target in slide_meta["targets"] if target["target_id"] == title.shape_id)
    assert body_target["text_category"] == "instructional_placeholder"
    assert body_target["replacement_mode"] == "replace" and body_target["clear_before_write"]
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name=slide_meta["layout_name"], targets=[
            TargetContent(target_kind="shape", target_id=title_target["target_id"], content_type="title", text="Evidence Overview"),
            TargetContent(target_kind="shape", target_id=body_target["target_id"], content_type="bullets", bullets=["Evidence supports practical operational decisions"]),
        ],
    )])
    output, _ = populate_presentation(source, content, tmp_path / "generated", metadata)
    generated = Presentation(output)
    all_text = "\n".join(shape.text for shape in generated.slides[0].shapes if getattr(shape, "has_text_frame", False))
    assert "Describe the supporting" not in all_text
    assert "Insert an optional" not in all_text
    assert "Example Corp" in all_text
    assert not [issue for issue in validate_presentation_layout(output, metadata)
                if issue.category == "unresolved_instruction"]


def test_unresolved_instruction_is_detected(tmp_path):
    source = tmp_path / "remaining.pptx"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    shape.text = "Explain the proposed solution here"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    deck.save(source)
    metadata = parse_template(source, file_path="templates/remaining.pptx")
    issues = validate_presentation_layout(source, metadata)
    assert any(issue.category == "unresolved_instruction" for issue in issues)


def test_professional_template_offline_manual_smoke(tmp_path):
    source = Path(__file__).resolve().parents[1] / "templates" / "Professional_PPT_Template.pptx"
    if not source.is_file():
        pytest.skip("Professional template is not present in this checkout")
    metadata = parse_template(source, file_path=source.name, template_name="Professional PPT Template")
    data = source.read_bytes()
    metadata.update({
        "checksum_sha256": checksum_sha256(data), "minio_bucket": "mock-templates",
        "minio_object_key": "professional.pptx", "pptx_binary_bytes": data,
    })
    class Retriever:
        def get_full(self, template_id): return deepcopy(metadata)
    class Watsonx:
        def chat(self, model_id, messages, **kwargs):
            payload = json.loads(messages[1]["content"].split("\n", 1)[1])
            slide = payload["planned_slide"]
            count = int(slide["generation_limits"]["requested_bullet_count"])
            bullets = [
                f"Source evidence supports section {slide['slide_number']} point {index}"
                for index in range(1, count + 1)
            ]
            return SimpleNamespace(text=json.dumps({"slides": [{
                "slide_number": slide["slide_number"], "section_type": slide["section_type"],
                "title": slide.get("section_title") or slide["section_type"].replace("_", " ").title(),
                "bullets": bullets, "speaker_notes": None,
            }]}))
    result = GenerationPipeline(Retriever(), Watsonx(), source.parent, tmp_path / "generated").run_manual(
        template_id=metadata["template_id"], model_id="mock", topic="Professional Research",
        complete_demand="Create a grounded professional presentation.",
        source_content="Source evidence supports professional research, planning, methodology, findings, and conclusions.",
        audience="Stakeholders", tone="Professional", desired_slide_count=14,
        required_sections=[], visual_preferences="", additional_instructions="",
    )
    generated = Presentation(result.output_path)
    visible = "\n".join(shape.text for slide in generated.slides for shape in slide.shapes
                         if getattr(shape, "has_text_frame", False))
    assert len(generated.slides) == 14
    assert not INSTRUCTION_TEXT_RE.search(visible)
    assert all(len(" ".join(target.content.split())) > 1 for slide in result.content.slides
               for target in slide.targets if target.content_type == "title")
