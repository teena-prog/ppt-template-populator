import json
from types import SimpleNamespace

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from src.pipeline import GenerationPipeline
from src.ppt_populator import populate_presentation
from src.prompt_builder import compact_template
from src.template_binary import checksum_sha256, temporary_pptx
from src.template_parser import parse_template
from src.target_metadata import required_targets_by_slide


def _text(slide, text, left, top, width, height, size=18):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    run = shape.text_frame.paragraphs[0].add_run(); run.text = text; run.font.size = Pt(size)
    return shape


def _fill_until(slide, target_id):
    while max((shape.shape_id for shape in slide.shapes), default=1) < target_id - 1:
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.1), Inches(.1), Inches(.05), Inches(.05))


def _realistic_canva(path):
    prs = Presentation()
    for number in range(1, 5):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _text(slide, f"Unique slide {number} title", .8, .4, 8, .7, 30)
        if number == 3:
            _fill_until(slide, 12)
            _text(slide, "First substantial column content", .8, 2, 4, 2, 16)
            _text(slide, "Second substantial column content", 5, 2, 4, 2, 16)
        elif number == 4:
            _fill_until(slide, 24)
            _text(slide, "First substantial metrics column", .5, 2, 2.8, 2, 16)
            _text(slide, "Second substantial metrics column", 3.5, 2, 2.8, 2, 16)
            _text(slide, "Third substantial metrics column", 6.5, 2, 2.8, 2, 16)
        _text(slide, f"Optional caption {number}", .8, 6, 2, .3, 10)
        _text(slide, "Cloudlet Tech", 7.5, 6, 1.5, .3, 10)
    prs.save(path)


class RecoveryWatson:
    def __init__(self): self.recovery_request = None
    def chat(self, model_id, messages, **kwargs):
        payload = json.loads(messages[1]["content"].split("\n", 1)[-1])
        if "ordered_targets" in payload:
            self.recovery_request = payload["ordered_targets"]
            return SimpleNamespace(text=json.dumps({"contents": [f"Recovered {item['target_id']}" for item in self.recovery_request]}))
        slides = []
        for slide in payload["template_structure"]:
            targets = [{"target_kind": item["target_kind"], "target_id": item["target_id"], "content": f"Generated {item['target_id']}"} for item in slide["targets"] if (slide["slide_number"], item["target_id"]) not in {(3, 12), (4, 24)}]
            slides.append({"slide_number": slide["slide_number"], "layout_name": slide["layout_name"], "targets": targets})
        return SimpleNamespace(text=json.dumps({"slides": slides}))


def test_equal_columns_brand_protection_focused_recovery_and_pptx(tmp_path):
    source = tmp_path / "canva-columns.pptx"; _realistic_canva(source)
    metadata = parse_template(source, file_path="templates/canva-columns.pptx")
    required = required_targets_by_slide(metadata)
    assert {12, 13}.issubset({target_id for kind, target_id in required[3]})
    assert {24, 25, 26}.issubset({target_id for kind, target_id in required[4]})
    normalized_targets = [target for slide in metadata["slides"] for target in slide["targets"]]
    brands = [target for target in normalized_targets if target["existing_text"] == "Cloudlet Tech"]
    assert brands and all(not target["replaceable"] and not target["required"] for target in brands)
    prompt = compact_template(metadata)
    prompt_keys = [(slide["slide_number"], target["target_id"]) for slide in prompt for target in slide["targets"]]
    assert len(prompt_keys) == len(set(prompt_keys))
    assert all("Optional caption" not in target["existing_text"] for slide in prompt for target in slide["targets"])

    watson = RecoveryWatson(); pipeline = GenerationPipeline(None, watson, tmp_path, tmp_path / "generated")
    content, warnings, calls = pipeline._generate_validated(model_id="model", template=metadata, message_args={"topic": "Startup", "source_content": "Facts", "audience": "Investors", "tone": "Clear", "additional_instructions": ""})
    assert calls == 2
    assert {(item["slide_number"], item["target_id"]) for item in watson.recovery_request} == {(3, 12), (4, 24)}
    generated = {(slide.slide_number, target.target_id): target.content for slide in content.slides for target in slide.targets}
    assert generated[(3, 13)] == "Generated 13" and generated[(4, 25)] == "Generated 25"
    assert generated[(3, 12)] == "Recovered 12" and generated[(4, 24)] == "Recovered 24"
    data = source.read_bytes()
    with temporary_pptx(data, checksum_sha256(data)) as temporary_source:
        temporary_name = temporary_source
        output, _ = populate_presentation(temporary_source, content, tmp_path / "generated", metadata)
    assert not temporary_name.exists() and output.exists()
    reopened = Presentation(output)
    assert len(reopened.slides) == 4 and any("Focused recovery" in warning for warning in warnings)


class FailedRecoveryWatson(RecoveryWatson):
    def chat(self, model_id, messages, **kwargs):
        response = super().chat(model_id, messages, **kwargs)
        if self.recovery_request is not None: return SimpleNamespace(text=json.dumps({"contents": []}))
        return response


def test_focused_recovery_failure_falls_back_to_placeholder_content(tmp_path):
    """When even the retried recovery calls omit a required target, the pipeline
    must still finish (with a placeholder + warning) instead of failing the whole
    presentation — see the 'Unresolved required target keys' UI regression."""
    source = tmp_path / "canva-columns.pptx"; _realistic_canva(source)
    metadata = parse_template(source, file_path="templates/canva-columns.pptx")
    pipeline = GenerationPipeline(None, FailedRecoveryWatson(), tmp_path, tmp_path)
    content, warnings, calls = pipeline._generate_validated(model_id="model", template=metadata, message_args={"topic": "Startup", "source_content": "Facts", "audience": "Investors", "tone": "Clear", "additional_instructions": ""})
    generated = {(slide.slide_number, target.target_id): target.content for slide in content.slides for target in slide.targets}
    assert (3, 12) in generated and generated[(3, 12)]
    assert (4, 24) in generated and generated[(4, 24)]
    fallback_warning = next(warning for warning in warnings if "could not be generated by the model" in warning)
    assert "(3, 'shape', 12)" in fallback_warning and "(4, 'shape', 24)" in fallback_warning


def test_target_budget_chunking_keeps_large_slide_whole(tmp_path):
    source = tmp_path / "canva-columns.pptx"; _realistic_canva(source)
    metadata = parse_template(source, file_path="templates/canva-columns.pptx")
    pipeline = GenerationPipeline(None, None, tmp_path, tmp_path, max_required_targets_per_chunk=2)
    chunks = pipeline._plan_chunks(metadata)
    assert all(len(chunk) <= 4 for chunk in chunks)
    assert [4] in chunks
