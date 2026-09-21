from src.document_extractor import DocumentExtractionError, ExtractedDocument
from src.source_batch import SourceUpload, process_source_batch
from src.ui_state import sources_ready
from src.slide_planner import build_factual_summary
from src.slide_planner import build_slide_plan
from src.template_parser import parse_template
from src.semantic_content import map_semantic_to_targets
from src.response_models import SemanticSlideContent, SemanticSlideResponse
from src.ppt_populator import populate_presentation
from pptx import Presentation

MB = 1024 * 1024

def batch(items, monkeypatch, extractor=None, per_file=1, total=3):
    extractor = extractor or (lambda data, name: ExtractedDocument(name.rsplit(".", 1)[0], data.decode()))
    monkeypatch.setattr("src.source_batch.extract_text", extractor)
    return process_source_batch(items, max_file_bytes=per_file * MB, max_total_bytes=total * MB)

def test_one_valid_file_preserves_single_file_workflow(monkeypatch):
    result = batch([SourceUpload("one.txt", b"One grounded paragraph.", "text/plain")], monkeypatch)
    assert len(result.valid_records) == 1 and result.title == "one"
    assert result.combined_text == "[SOURCE 1: one.txt]\nOne grounded paragraph."

def test_multiple_mixed_files_form_one_combined_source(monkeypatch):
    items = [SourceUpload("a.pdf", b"PDF facts", "application/pdf"), SourceUpload("b.docx", b"DOCX facts", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), SourceUpload("c.txt", b"TXT facts", "text/plain")]
    result = batch(items, monkeypatch)
    assert len(result.valid_records) == 3
    assert result.combined_text.count("[SOURCE ") == 3
    assert all(name in result.combined_text for name in ("a.pdf", "b.docx", "c.txt"))

def test_duplicate_checksum_is_not_extracted_twice(monkeypatch):
    result = batch([SourceUpload("a.txt", b"same", "text/plain"), SourceUpload("b.txt", b"same", "text/plain")], monkeypatch)
    assert len(result.valid_records) == 1
    assert result.records[1].duplicate and result.records[1].extraction_status == "duplicate"

def test_invalid_file_does_not_stop_valid_file(monkeypatch):
    def extractor(data, name):
        if name == "bad.pdf": raise DocumentExtractionError("Unreadable document.")
        return ExtractedDocument("Good", "Grounded content")
    result = batch([SourceUpload("bad.pdf", b"bad", "application/pdf"), SourceUpload("good.txt", b"good", "text/plain")], monkeypatch, extractor)
    assert len(result.valid_records) == 1 and result.records[0].extraction_status == "failed"

def test_all_invalid_and_empty_file(monkeypatch):
    result = batch([SourceUpload("empty.txt", b"", "text/plain"), SourceUpload("bad.exe", b"x", "application/octet-stream")], monkeypatch)
    assert not result.valid_records
    assert [record.extraction_status for record in result.records] == ["invalid", "invalid"]

def test_per_file_and_total_size_limits(monkeypatch):
    per_file = process_source_batch([SourceUpload("large.txt", b"x" * 11, "text/plain")], max_file_bytes=10, max_total_bytes=20)
    assert not per_file.valid_records and "per-file" in per_file.records[0].error
    total = process_source_batch([SourceUpload("a.txt", b"x" * 6, "text/plain"), SourceUpload("b.txt", b"y" * 6, "text/plain")], max_file_bytes=10, max_total_bytes=11)
    assert total.batch_error and len(total.records) == 2
    assert all(record.extraction_status == "invalid" for record in total.records)

def test_source_boundaries_and_cross_file_paragraph_deduplication(monkeypatch):
    result = batch([SourceUpload("a.txt", b"Repeated fact.\n\nUnique A.", "text/plain"), SourceUpload("b.txt", b"Repeated fact.\n\nUnique B.", "text/plain")], monkeypatch)
    assert result.combined_text.count("Repeated fact.") == 1
    assert "[SOURCE 1: a.txt]" in result.combined_text and "[SOURCE 2: b.txt]" in result.combined_text
    assert "Unique A." in result.combined_text and "Unique B." in result.combined_text
    assert "[SOURCE" not in build_factual_summary(result.combined_text)

def test_mime_mismatch_is_rejected(monkeypatch):
    result = batch([SourceUpload("report.pdf", b"data", "text/plain")], monkeypatch)
    assert not result.valid_records and "does not match" in result.records[0].error

def test_unrelated_topic_requires_explicit_confirmation(monkeypatch):
    result = batch([SourceUpload("finance.txt", b"Revenue margin earnings forecast", "text/plain"), SourceUpload("biology.txt", b"Cellular genome protein enzyme", "text/plain")], monkeypatch)
    assert result.warnings
    assert not sources_ready(result, False)
    assert sources_ready(result, True)

def test_related_documents_do_not_require_confirmation(monkeypatch):
    result = batch([SourceUpload("one.txt", b"Climate emissions policy analysis", "text/plain"), SourceUpload("two.txt", b"Emissions targets support climate planning", "text/plain")], monkeypatch)
    assert not result.warnings and sources_ready(result, False)

def test_multiple_sources_build_one_summary_plan_and_ppt(monkeypatch, sample_pptx, tmp_path):
    result = batch([
        SourceUpload("problem.txt", b"Evidence identifies a practical problem. Teams define clear objectives.", "text/plain"),
        SourceUpload("solution.txt", b"Evidence supports a practical solution. Findings guide validation and conclusions.", "text/plain"),
    ], monkeypatch)
    summary = build_factual_summary(result.combined_text)
    assert "[SOURCE" not in summary
    template = parse_template(sample_pptx, file_path="templates/sample.pptx")
    planned = build_slide_plan(template, result.combined_text, 10)
    semantic = SemanticSlideResponse(slides=[SemanticSlideContent(
        slide_number=slide["slide_number"], section_type=slide["section_type"],
        title=result.title if slide["slide_number"] == 1 else slide["planned_section_title"],
        bullets=[] if slide["slide_number"] in {1, 10} else ["Evidence supports practical decisions"],
    ) for slide in planned["slides"]])
    content, _ = map_semantic_to_targets(semantic, planned, result.title)
    output, _ = populate_presentation(sample_pptx, content, tmp_path / "generated", planned)
    assert output.exists() and len(Presentation(output).slides) == 10
    visible = " ".join(shape.text for slide in Presentation(output).slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert "[SOURCE" not in visible
