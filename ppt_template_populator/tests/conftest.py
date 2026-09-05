from __future__ import annotations

import sys
from pathlib import Path
import pytest
from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Original title"
    slide.placeholders[1].text = "Original subtitle"
    path = tmp_path / "sample.pptx"
    prs.save(path)
    return path


@pytest.fixture
def template_metadata(sample_pptx: Path) -> dict:
    from src.template_parser import parse_template
    return parse_template(sample_pptx, file_path="templates/sample.pptx", template_id="template-1")


@pytest.fixture
def indexed_template(template_metadata: dict, sample_pptx: Path) -> dict:
    from src.template_binary import checksum_sha256
    from src.template_profile import build_template_profile
    data = sample_pptx.read_bytes()
    template_metadata.update({"use_cases": ["pitch"], "target_audiences": ["investors"], "tones": ["professional"], "visual_style": "minimal", "supported_sections": ["problem", "solution"], "filename": "sample.pptx", "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "file_size": len(data), "checksum_sha256": checksum_sha256(data), "template_embedding": [0.1, 0.2, 0.3], "embedding_model_id": "ibm/test-embedding", "embedding_dimensions": 3, "minio_bucket": "ppt-templates", "minio_object_key": "template-1.pptx"})
    template_metadata["template_profile"] = build_template_profile(template_metadata)
    template_metadata["pptx_binary_bytes"] = data
    return template_metadata


@pytest.fixture
def valid_payload(template_metadata: dict) -> dict:
    slide = template_metadata["slides"][0]
    targets = [{"target_kind": "placeholder", "target_id": shape["placeholder_idx"], "content": f"Content {shape['placeholder_idx']}"} for shape in slide["shapes"] if shape["is_placeholder"] and shape["has_text_frame"]]
    return {"presentation_title": "Test", "slides": [{"slide_number": 1, "layout_name": slide["layout_name"], "targets": targets, "speaker_notes": None}]}
