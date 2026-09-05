"""The house typography standard must be enforced on every populated target
regardless of what the source template originally used: Aptos throughout,
fixed per-role sizes/weights, a consistent line spacing, sub-bullets
(outline level >= 1 inside a body target) one step down from body text, and
PowerPoint's own dynamic shrink-to-fit disabled so those sizes hold."""
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from src.ppt_populator import populate_presentation
from src.response_models import PresentationContent, SlideContent, TargetContent


def _build_source_pptx(path: Path) -> tuple[int, int]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content

    title = slide.shapes.title
    title.text = "Original title"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].runs[0].font.name = "Calibri"

    body = slide.placeholders[1]
    text_frame = body.text_frame
    text_frame.text = "Original bullet"
    text_frame.paragraphs[0].level = 0
    text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    sub = text_frame.add_paragraph()
    sub.text = "Original sub bullet"
    sub.level = 1
    sub.runs[0].font.size = Pt(22)

    prs.save(path)
    return title.placeholder_format.idx, body.placeholder_format.idx


def _template(title_idx: int, body_idx: int) -> dict:
    def target(idx, role, text, limit=200):
        return {
            "slide_number": 1, "target_kind": "placeholder", "target_id": idx, "role": role,
            "existing_text": text, "approximate_max_characters": limit, "max_content_length": limit,
            "required": True, "replaceable": True, "position": {"left": 0, "top": 0, "width": 100, "height": 40},
        }
    return {
        "template_id": "typography-check", "slide_count": 1,
        "slides": [{
            "slide_number": 1, "layout_name": "Title and Content",
            "targets": [target(title_idx, "title", "Original title"), target(body_idx, "body", "Original bullet")],
        }],
    }


def test_standard_typography_overrides_original_font_by_role_and_level(tmp_path):
    source = tmp_path / "source.pptx"
    title_idx, body_idx = _build_source_pptx(source)
    template = _template(title_idx, body_idx)

    content = PresentationContent(presentation_title="Test", slides=[
        SlideContent(slide_number=1, layout_name="Title and Content", targets=[
            TargetContent(target_kind="placeholder", target_id=title_idx, content="New Title"),
            TargetContent(target_kind="placeholder", target_id=body_idx, content="New top bullet\nNew sub bullet"),
        ]),
    ])

    output, warnings = populate_presentation(source, content, tmp_path / "generated", template)
    assert warnings == []

    result = Presentation(output)
    slide = result.slides[0]
    title_shape = next(s for s in slide.shapes if s.placeholder_format is not None and s.placeholder_format.idx == title_idx)
    body_shape = next(s for s in slide.shapes if s.placeholder_format is not None and s.placeholder_format.idx == body_idx)

    title_run = title_shape.text_frame.paragraphs[0].runs[0]
    assert title_run.font.name == "Aptos" and title_run.font.size == Pt(30) and title_run.font.bold is True

    body_paragraphs = body_shape.text_frame.paragraphs
    assert body_paragraphs[0].text == "New top bullet" and body_paragraphs[0].level == 0
    top_run = body_paragraphs[0].runs[0]
    assert top_run.font.name == "Aptos" and top_run.font.size == Pt(20) and top_run.font.bold is False

    # Second line reused the original sub-bullet paragraph's level-1 pPr, so
    # it renders one step down at the sub-bullet size.
    assert body_paragraphs[1].text == "New sub bullet" and body_paragraphs[1].level == 1
    sub_run = body_paragraphs[1].runs[0]
    assert sub_run.font.name == "Aptos" and sub_run.font.size == Pt(18) and sub_run.font.bold is False

    assert body_shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert body_shape.text_frame.word_wrap is True
