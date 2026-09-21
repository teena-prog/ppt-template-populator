import io
from types import SimpleNamespace

import pytest
from pptx import Presentation

import src.template_catalogue as catalogue
from src.template_binary import checksum_sha256


class FakeES:
    def __init__(self, docs=None):
        self.docs = dict(docs or {})
        self.deleted = []

    def search(self, *, index, query, size, _source=None):
        hits = [{"_id": key, "_source": value} for key, value in self.docs.items()]
        return {"hits": {"total": {"value": len(hits)}, "hits": hits[:size]}}

    def delete(self, *, index, id, refresh):
        self.deleted.append((index, id))
        return {"result": "deleted" if self.docs.pop(id, None) is not None else "not_found"}


class FakeMinio:
    def __init__(self, buckets=None):
        self.buckets = {name: dict(objects) for name, objects in (buckets or {}).items()}
        self.removed = []

    def list_objects(self, bucket, recursive=True):
        return [SimpleNamespace(object_name=key) for key in self.buckets.get(bucket, {})]

    def remove_object(self, bucket, key):
        self.removed.append((bucket, key)); self.buckets[bucket].pop(key)

    def bucket_exists(self, bucket): return bucket in self.buckets
    def make_bucket(self, bucket): self.buckets[bucket] = {}
    def put_object(self, bucket, key, stream, length, content_type): self.buckets[bucket][key] = stream.read()
    def get_object(self, bucket, key):
        data = self.buckets[bucket][key]
        return SimpleNamespace(read=lambda: data, close=lambda: None, release_conn=lambda: None)


def _pptx(slides=17):
    prs = Presentation()
    while len(prs.slides) < slides:
        prs.slides.add_slide(prs.slide_layouts[6])
    output = io.BytesIO(); prs.save(output); return output.getvalue()


def _doc(data):
    return {"template_id": "new-template", "template_name": "Professional", "filename": "Professional_PPT_Template.pptx", "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "checksum_sha256": checksum_sha256(data), "slide_count": 17, "slides": [{"slide_number": number + 1} for number in range(17)], "template_embedding": [0.1], "embedding_model_id": "mock", "embedding_dimensions": 1, "minio_bucket": "ppt-templates", "minio_object_key": "new-template.pptx"}


def test_scoped_deletion_preserves_unrelated_bucket():
    es = FakeES({"old": {"template_id": "old", "minio_bucket": "ppt-templates", "minio_object_key": "old.pptx"}})
    minio = FakeMinio({"ppt-templates": {"old.pptx": b"old"}, "uploads": {"source.pdf": b"protected"}})
    manifest = catalogue.build_manifest(es, "ppt-templates")
    assert catalogue.clear_catalogue(es, minio, "ppt-templates", manifest) == (1, 1)
    assert minio.buckets["uploads"] == {"source.pdf": b"protected"}
    assert es.deleted == [("ppt_templates", "old")]


def test_unidentifiable_template_bucket_object_blocks_all_deletion():
    es = FakeES({"old": {"template_id": "old", "minio_bucket": "ppt-templates", "minio_object_key": "old.pptx"}})
    minio = FakeMinio({"ppt-templates": {"old.pptx": b"old", "notes.txt": b"unknown"}})
    with pytest.raises(catalogue.CatalogueReplacementError, match="nothing was deleted"):
        catalogue.clear_catalogue(es, minio, "ppt-templates", catalogue.build_manifest(es, "ppt-templates"))
    assert "old" in es.docs and "old.pptx" in minio.buckets["ppt-templates"]


def test_successful_replacement_and_full_checksum_retrieval(monkeypatch, tmp_path):
    data = _pptx(); document = _doc(data)
    es = FakeES({"old": {"template_id": "old", "minio_bucket": "ppt-templates", "minio_object_key": "old.pptx"}})
    minio = FakeMinio({"ppt-templates": {"old.pptx": b"old"}, "uploads": {"keep": b"safe"}})
    class Indexer:
        def __init__(self, client): self.client = client
        def index(self, doc): self.client.docs[doc["template_id"]] = dict(doc)
    monkeypatch.setattr(catalogue, "TemplateIndexer", Indexer)
    result = catalogue.replace_catalogue(es, minio, "ppt-templates", document, data, tmp_path / "manifest.json")
    assert (result.deleted_documents, result.deleted_objects, result.slide_count) == (1, 1, 17)
    assert list(es.docs) == ["new-template"] and list(minio.buckets["ppt-templates"]) == ["new-template.pptx"]
    assert checksum_sha256(minio.buckets["ppt-templates"]["new-template.pptx"]) == document["checksum_sha256"]
    assert "template_embedding" not in (tmp_path / "manifest.json").read_text()


def test_index_failure_rolls_back_new_object(monkeypatch):
    data = _pptx(); document = _doc(data); es = FakeES(); minio = FakeMinio({"ppt-templates": {}})
    class FailingIndexer:
        def __init__(self, client): pass
        def index(self, doc): raise RuntimeError("index failed")
    monkeypatch.setattr(catalogue, "TemplateIndexer", FailingIndexer)
    with pytest.raises(catalogue.CatalogueReplacementError, match="rolled back"):
        catalogue.write_replacement(es, minio, document, data)
    assert minio.buckets["ppt-templates"] == {}
