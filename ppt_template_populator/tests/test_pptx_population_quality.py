from __future__ import annotations

from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR_INDEX
from pptx.util import Inches, Pt

from src.ppt_populator import _contrast, _fit_content_to_capacity, populate_presentation
from src.response_models import PresentationContent, SemanticSlideContent, SemanticSlideResponse, SlideContent, TargetContent
from src.semantic_content import map_semantic_to_targets
from src.template_binary import duplicate_package_members
from src.template_parser import parse_template


def _quality_template(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.6), Inches(.35), Inches(11.8), Inches(.65))
    title.text = "Section Heading"; title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    strip = slide.shapes.add_textbox(Inches(.8), Inches(1.25), Inches(11.5), Inches(.4))
    strip.text = "Static supporting label"; strip.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    body = slide.shapes.add_textbox(Inches(1), Inches(1.9), Inches(11.1), Inches(4.9))
    body.text = "Editable body content for this section"; body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    body.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(20, 40, 80)
    prs.save(path)
    return title.shape_id, strip.shape_id, body.shape_id


def test_short_strip_is_not_a_body_target_and_largest_body_is_used(tmp_path):
    source = tmp_path / "quality.pptx"
    _, strip_id, body_id = _quality_template(source)
    metadata = parse_template(source, file_path="templates/quality.pptx")
    slide = metadata["slides"][0]
    strip = next(shape for shape in slide["shapes"] if shape["shape_id"] == strip_id)
    assert strip["role"] == "short_label" and not strip["is_writable_target"]
    assert strip["capacity"]["maximum_lines"] == 1
    assert slide["body_targets"] == [{"target_kind": "shape", "target_id": body_id}]

    slide["section_type"] = "introduction"
    for target in slide["targets"]:
        target["required"] = target["target_id"] in {slide["title_target"]["target_id"], body_id}
        target["source_excerpt"] = "Grounded evidence supports the presentation."
    semantic = SemanticSlideResponse(slides=[SemanticSlideContent(
        slide_number=1, section_type="introduction", title="Introduction",
        bullets=["Grounded evidence supports the presentation"],
    )])
    mapped, _ = map_semantic_to_targets(semantic, metadata, "Document title")
    assert [(target.target_id, target.content_type) for target in mapped.slides[0].targets] == [
        (slide["title_target"]["target_id"], "title"), (body_id, "bullets")
    ]


def test_geometry_fitting_shortens_at_word_boundaries():
    target = {"target_kind": "shape", "target_id": 9, "role": "body",
              "capacity": {"characters_per_line": 20, "maximum_lines": 2},
              "approximate_max_characters": 40}
    fitted, changed = _fit_content_to_capacity(
        "This grounded sentence contains too many words for the available two line destination", target
    )
    assert changed and len(fitted.split()) < 13 and not fitted.endswith(" ")


def test_safe_output_preserves_static_text_and_explicit_font_colour(tmp_path):
    source = tmp_path / "quality.pptx"
    title_id, strip_id, body_id = _quality_template(source)
    metadata = parse_template(source, file_path="templates/quality.pptx")
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name="Blank", targets=[
            TargetContent(target_kind="shape", target_id=title_id, content_type="title", text="Introduction"),
            TargetContent(target_kind="shape", target_id=body_id, content_type="bullets", bullets=["Grounded point"]),
        ],
    )])
    output, _ = populate_presentation(source, content, tmp_path / "generated", metadata)
    assert duplicate_package_members(output) == []
    with ZipFile(output) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    reopened = Presentation(output)
    shapes = {shape.shape_id: shape for shape in reopened.slides[0].shapes}
    assert shapes[strip_id].text == "Static supporting label"
    assert shapes[body_id].text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(20, 40, 80)


def test_inherited_theme_text_gets_explicit_readable_colour(tmp_path):
    source = tmp_path / "theme.pptx"
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(5, 15, 25)
    shape.text = "Theme text"
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(20); run.font.color.theme_color = MSO_THEME_COLOR_INDEX.TEXT_1
    prs.save(source)
    metadata = {"slide_count": 1, "slides": [{"slide_number": 1, "layout_name": "Blank", "targets": [{
        "slide_number": 1, "target_kind": "shape", "target_id": shape.shape_id, "role": "body",
        "existing_text": "Theme text", "replaceable": True, "required": True,
        "approximate_max_characters": 200, "max_content_length": 200,
        "capacity": {"characters_per_line": 50, "maximum_lines": 5},
    }]}]}
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name="Blank", targets=[TargetContent(
            target_kind="shape", target_id=shape.shape_id, content_type="bullets", bullets=["Readable point"]
        )],
    )])
    output, _ = populate_presentation(source, content, tmp_path / "generated", metadata)
    result = Presentation(output); generated = result.slides[0].shapes[0]
    rgb = tuple(generated.text_frame.paragraphs[0].runs[0].font.color.rgb)
    assert _contrast(rgb, (5, 15, 25)) >= 3.0


def test_failed_deck_fixture_and_professional_template_regression():
    root = __import__("pathlib").Path(__file__).resolve().parents[1]
    failed = root / "generated" / "presentation_f2a850c5f6bb4fd8805d39ac20d3b706.pptx"
    template = root / "templates" / "Professional_PPT_Template.pptx"
    if not failed.exists() or not template.exists():
        __import__("pytest").skip("Local regression artifacts are not available")
    assert duplicate_package_members(failed) == []
    Presentation(failed)
    metadata = parse_template(template, file_path="templates/Professional_PPT_Template.pptx")
    content_slides = [slide for slide in metadata["slides"] if slide["slide_number"] in range(3, 9)]
    for slide in content_slides:
        strip = next(shape for shape in slide["shapes"] if 0.35 <= shape["position"]["height"] / 914400 <= 0.45)
        assert strip["role"] == "short_label" and not strip["is_writable_target"]
        body_ids = {target["target_id"] for target in slide["targets"] if target["replaceable"] and target["role"] in {"body", "bullets"}}
        assert strip["shape_id"] not in body_ids and body_ids
