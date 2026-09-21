from pptx import Presentation
from pptx.util import Inches, Pt

from src.content_normalizer import normalize_bullets
from src.ppt_populator import _fit_content_to_capacity, populate_presentation
from src.response_models import PresentationContent, SlideContent, TargetContent
from src.template_parser import parse_template


def test_wrapped_pdf_lines_become_one_complete_sentence():
    value = (
        "Start the local model - a 120-billion-parameter model\n"
        "running on our own GPU, from weights on\n"
        "our own disk, with no internet involved."
    )
    assert normalize_bullets(value) == [
        "Start the local model - a 120-billion-parameter model running on our own GPU, from weights on our own disk, with no internet involved."
    ]


def test_standalone_number_is_removed_and_following_wraps_are_joined():
    value = (
        "2.\nMove summarisation off the cloud API -\n"
        "nothing about a railway conversation should be sent to\n"
        "an external service."
    )
    assert normalize_bullets(value) == [
        "Move summarisation off the cloud API - nothing about a railway conversation should be sent to an external service."
    ]


def test_reliable_markers_remain_separate_bullets():
    assert normalize_bullets("• First complete point\n• Second complete point") == [
        "First complete point", "Second complete point"
    ]
    assert normalize_bullets("1. First numbered point\n2) Second numbered point") == [
        "First numbered point", "Second numbered point"
    ]


def test_complete_json_list_items_are_not_merged():
    items = ["First complete grounded sentence.", "Second complete grounded sentence."]
    assert normalize_bullets(items) == items


def test_meaningful_lowercase_bullet_is_capitalized_without_repair():
    assert normalize_bullets([
        "Local inference protects sensitive conversations.",
        "using local infrastructure can also reduce external dependency.",
    ]) == [
        "Local inference protects sensitive conversations.",
        "Using local infrastructure can also reduce external dependency.",
    ]


def test_clear_wrapped_continuation_is_merged_using_multiple_signals():
    assert normalize_bullets([
        "Start the local model from weights stored",
        "on the local disk without internet access.",
    ]) == ["Start the local model from weights stored on the local disk without internet access."]


def test_intentional_lowercase_technical_forms_are_preserved():
    values = [
        "e.g. local deployment", "i.e. no external service", "pH measurement",
        "iOS application", "e-commerce workflow", "https://example.test", "model_name output",
    ]
    assert normalize_bullets(values) == values


def test_concise_noun_phrase_bullets_are_accepted():
    values = [
        "Reduced dependence on external cloud services",
        "Local processing for sensitive railway conversations",
        "Faster response during connectivity disruptions",
    ]
    assert normalize_bullets(values) == values


def test_structured_standalone_marker_remains_invalid():
    try:
        normalize_bullets(["2.", "A complete point"])
    except ValueError as exc:
        assert "standalone list markers" in str(exc)
    else:
        raise AssertionError("A standalone structured marker must remain invalid")


def _visual_instruction_template(path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(8), Inches(.7))
    title.text = "Section Title"; title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    body = slide.shapes.add_textbox(Inches(.8), Inches(1.5), Inches(7.5), Inches(3.5))
    body.text = "Replace with body content"; body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    visual = slide.shapes.add_textbox(Inches(9), Inches(1.5), Inches(3), Inches(1))
    visual.text = "VISUAL"
    visual.fill.solid()
    border_xml = visual._element.spPr.xml
    instruction = slide.shapes.add_textbox(Inches(9), Inches(2.7), Inches(3), Inches(1))
    instruction.text = "Replace with relevant visual"
    brand = slide.shapes.add_textbox(Inches(.5), Inches(7), Inches(2), Inches(.2))
    brand.text = "BRAND LEGAL"; brand.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    deck.save(path)
    return title.shape_id, body.shape_id, visual.shape_id, instruction.shape_id, border_xml


def test_complete_bullet_is_one_paragraph_and_visual_instructions_are_cleared(tmp_path):
    source = tmp_path / "visual.pptx"
    title_id, body_id, visual_id, instruction_id, _ = _visual_instruction_template(source)
    template = parse_template(source, file_path="templates/visual.pptx")
    generated = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name="Blank", targets=[
            TargetContent(target_kind="shape", target_id=title_id, content_type="title", text="Complete Title"),
            TargetContent(target_kind="shape", target_id=body_id, content_type="bullets", bullets=[
                "One complete sentence remains inside one PowerPoint bullet paragraph."
            ]),
        ],
    )])
    output, _ = populate_presentation(source, generated, tmp_path / "generated", template)
    result = Presentation(output)
    shapes = {shape.shape_id: shape for shape in result.slides[0].shapes}
    assert len([p for p in shapes[body_id].text_frame.paragraphs if p.text.strip()]) == 1
    assert shapes[visual_id].text == ""
    assert shapes[instruction_id].text == ""
    assert shapes[visual_id].fill.type is not None
    assert any(shape.text == "BRAND LEGAL" for shape in result.slides[0].shapes if shape.has_text_frame)


def test_long_bullet_is_shortened_as_one_item_not_fragmented():
    target = {
        "target_kind": "shape", "target_id": 42, "role": "body",
        "approximate_max_characters": 48,
        "capacity": {"characters_per_line": 24, "maximum_lines": 2},
    }
    fitted, changed = _fit_content_to_capacity(
        "This complete grounded sentence contains additional words that cannot fit inside the available body region.",
        target,
    )
    assert changed
    assert "\n" not in fitted
    assert len(fitted.split()) >= 3
