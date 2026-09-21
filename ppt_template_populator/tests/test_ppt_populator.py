from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from src.ppt_populator import populate_presentation
from src.response_validator import validate_response
from src.response_models import PresentationContent, SlideContent, TargetContent


def test_powerpoint_output_generation(sample_pptx: Path, template_metadata, valid_payload, tmp_path: Path):
    content, _ = validate_response(valid_payload, template_metadata)
    output, warnings = populate_presentation(sample_pptx, content, tmp_path / "generated")
    assert output.exists() and output.name.startswith("presentation_")
    prs = Presentation(output)
    for item in content.slides[0].placeholders:
        shape = next(s for s in prs.slides[0].shapes if s.is_placeholder and s.placeholder_format.idx == item.placeholder_id)
        assert shape.text == item.content
    assert warnings == []


def test_bullet_array_becomes_separate_powerpoint_bullet_paragraphs(sample_pptx, template_metadata, tmp_path):
    slide = template_metadata["slides"][0]
    title = next(target for target in slide["targets"] if target["role"] == "title")
    body = next(target for target in slide["targets"] if target["role"] == "subtitle")
    body["role"] = "body"
    payload = {"slides": [{"slide_number": 1, "layout_name": slide["layout_name"], "targets": [
        {"target_kind": title["target_kind"], "target_id": title["target_id"], "content_type": "title", "text": "Structured title"},
        {"target_kind": body["target_kind"], "target_id": body["target_id"], "content_type": "bullets", "bullets": ["First grounded point", "Second grounded point", "Third grounded point"]},
    ]}]}
    content, _ = validate_response(payload, template_metadata, presentation_title="Test")
    output, _ = populate_presentation(sample_pptx, content, tmp_path / "generated", template_metadata)
    prs = Presentation(output)
    shape = next(item for item in prs.slides[0].shapes if item.is_placeholder and item.placeholder_format.idx == body["target_id"])
    assert [paragraph.text for paragraph in shape.text_frame.paragraphs] == payload["slides"][0]["targets"][1]["bullets"]
    assert all(paragraph._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None for paragraph in shape.text_frame.paragraphs)
    assert all(len(paragraph._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")) == 1 for paragraph in shape.text_frame.paragraphs)


def test_near_white_generated_text_is_made_readable(sample_pptx, template_metadata, tmp_path):
    presentation = Presentation(sample_pptx)
    body_shape = presentation.slides[0].placeholders[1]
    body_shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    source = tmp_path / "white-text.pptx"
    presentation.save(source)
    slide = template_metadata["slides"][0]
    body = next(target for target in slide["targets"] if target["role"] == "subtitle")
    body["role"] = "body"
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name=slide["layout_name"], targets=[TargetContent(
            target_kind=body["target_kind"], target_id=body["target_id"],
            content_type="bullets", bullets=["Readable generated point"],
        )],
    )])
    output, _ = populate_presentation(source, content, tmp_path / "generated", template_metadata)
    generated = Presentation(output)
    run = generated.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor(0x18, 0x35, 0x50)




def test_speaker_notes_are_written_only_to_slide_notes(sample_pptx, template_metadata, tmp_path):
    slide = template_metadata["slides"][0]
    title = next(target for target in slide["targets"] if target["role"] == "title")
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name=slide["layout_name"],
        targets=[TargetContent(target_kind=title["target_kind"], target_id=title["target_id"], content_type="title", text="Visible title")],
        speaker_notes="Private presenter guidance",
    )])
    output, _ = populate_presentation(sample_pptx, content, tmp_path / "generated", template_metadata)
    generated = Presentation(output)
    assert generated.slides[0].notes_slide.notes_text_frame.text == "Private presenter guidance"
    visible = " ".join(shape.text for shape in generated.slides[0].shapes if getattr(shape, "has_text_frame", False))
    assert "Private presenter guidance" not in visible


def test_copied_page_number_is_rewritten_to_output_slide_number(sample_pptx, template_metadata, tmp_path):
    prs = Presentation(sample_pptx)
    footer = prs.slides[0].shapes.add_textbox(0, 0, 200000, 200000)
    footer.text = "99"
    source = tmp_path / "numbered.pptx"; prs.save(source)
    from src.template_parser import parse_template
    metadata = parse_template(source, file_path="numbered.pptx")
    title = next(target for target in metadata["slides"][0]["targets"] if target["role"] == "title")
    content = PresentationContent(presentation_title="Deck", slides=[SlideContent(
        slide_number=1, layout_name=metadata["slides"][0]["layout_name"],
        targets=[TargetContent(target_kind=title["target_kind"], target_id=title["target_id"], content_type="title", text="Document title")],
    )])
    output, _ = populate_presentation(source, content, tmp_path / "generated", metadata)
    generated = Presentation(output)
    assert any(shape.text == "1" for shape in generated.slides[0].shapes if getattr(shape, "has_text_frame", False))
