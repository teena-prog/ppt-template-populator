import argparse
import json
from types import SimpleNamespace
from unittest.mock import MagicMock
import pytest
from pptx import Presentation
from scripts import ingest_template
from src.embedding_client import EmbeddingResult
from src.pipeline import GenerationPipeline
from src.template_indexer import DuplicateTemplateError, TemplateIndexer


def ingestion_args(path):
    return argparse.Namespace(pptx_path=path, name="Startup Pitch", description="Investor deck", category="Business", use_cases="startup pitch,fundraising", audiences="investors", tones="professional,persuasive", visual_style="minimal", supported_sections="problem,solution", max_size_mb=10, update_existing=False)


def test_valid_pptx_ingestion_builds_profile_and_binary(monkeypatch, sample_pptx):
    settings = SimpleNamespace(max_upload_mb=25, watsonx_embedding_model_id="embed-model", minio_templates_bucket="ppt-templates")
    monkeypatch.setattr(ingest_template, "get_settings", lambda: settings)
    manager = MagicMock(); manager.get_embeddings_model_specs.return_value = {"resources": [{"model_id": "embed-model"}]}
    watsonx = SimpleNamespace(api_client=SimpleNamespace(foundation_models=manager))
    embedder = MagicMock(); embedder.embed.return_value = EmbeddingResult([.1, .2, .3], "embed-model")
    document, data = ingest_template.build_ingestion_document(ingestion_args(sample_pptx), embedding_client=embedder, watsonx=watsonx)
    assert data and document["minio_bucket"] == "ppt-templates" and document["minio_object_key"] and document["checksum_sha256"]
    assert document["embedding_dimensions"] == 3
    assert "Startup Pitch" in document["template_profile"] and "problem" in document["template_profile"]


def test_duplicate_template_prevention(indexed_template):
    client = MagicMock(); client.indices.exists.return_value = False; client.exists.return_value = True
    with pytest.raises(DuplicateTemplateError): TemplateIndexer(client).index(indexed_template)
    client.index.assert_not_called()


def test_update_existing_ingestion_preserves_template_id(indexed_template):
    client = MagicMock(); client.indices.exists.return_value = False; client.exists.return_value = False
    client.search.return_value = {"hits": {"hits": [{"_source": {"template_id": "existing-id"}}]}}
    result = TemplateIndexer(client).index(indexed_template, update_existing=True)
    assert result == "existing-id"
    assert client.index.call_args.kwargs["id"] == "existing-id"
    assert "op_type" not in client.index.call_args.kwargs


def test_update_existing_cli_option():
    args = ingest_template.parser().parse_args(["template.pptx", "--name", "Template", "--update-existing"])
    assert args.update_existing is True

def test_update_existing_preserves_minio_reference_and_identity():
    rebuilt = {
        "template_id": "new", "minio_bucket": "bucket", "minio_object_key": "new.pptx",
        "created_at": "new-time", "description": "", "category": "General", "use_cases": [],
        "target_audiences": [], "tones": [], "visual_style": "", "supported_sections": [],
    }
    existing = {
        "template_id": "kept", "minio_bucket": "templates", "minio_object_key": "kept.pptx",
        "created_at": "old-time", "description": "Existing description", "category": "Business",
        "use_cases": ["strategy"], "target_audiences": ["executives"], "tones": ["formal"],
        "visual_style": "corporate", "supported_sections": ["title", "conclusion"],
    }
    result = ingest_template.preserve_existing_identity_and_storage(rebuilt, existing)
    assert result["template_id"] == "kept"
    assert result["minio_bucket"] == "templates" and result["minio_object_key"] == "kept.pptx"
    assert result["created_at"] == "old-time"
    assert result["description"] == "Existing description"
    assert result["category"] == "Business"
    assert result["use_cases"] == ["strategy"]
    assert result["target_audiences"] == ["executives"]
    assert result["tones"] == ["formal"]
    assert result["visual_style"] == "corporate"
    assert result["supported_sections"] == ["title", "conclusion"]


class AutomaticRetriever:
    def __init__(self, template): self.template = template; self.binary_requested = False
    def embedding_configurations(self): return [{"model_id": self.template["embedding_model_id"], "dimensions": self.template["embedding_dimensions"], "count": 1}]
    def hybrid_search(self, **kwargs):
        safe = {key: value for key, value in self.template.items() if key not in {"pptx_binary_bytes", "template_embedding"}}
        safe.update({"retrieval_score": 2.0, "requirement_match_score": .5, "matched_requirements": ["audience"], "unmatched_requirements": ["tone"]})
        return [safe]
    def get_full(self, template_id): self.binary_requested = True; return self.template


def test_complete_automatic_pipeline(indexed_template, valid_payload, tmp_path):
    selection = {"selected_template_id": indexed_template["template_id"], "confidence": .5, "reason": "Best candidate", "matched_requirements": ["audience"], "unmatched_requirements": ["tone"]}
    calls = 0
    def respond(model_id, messages, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            return SimpleNamespace(text=json.dumps(selection))
        payload = json.loads(messages[1]["content"].split("\n", 1)[1])
        if "planned_slide" in payload:
            slide = payload["planned_slide"]
            return SimpleNamespace(text=json.dumps({"slides": [{
                "slide_number": slide["slide_number"], "section_type": slide["section_type"],
                "title": slide.get("section_title") or "Grounded title",
                "bullets": ([] if slide["section_type"] in {"title", "qa"} else [
                    "Facts establish the central context", "Facts clarify the relevant evidence",
                    "Facts support the practical conclusion",
                ]), "speaker_notes": None,
            }]}))
        if "ordered_targets" in payload:
            values = []
            for target in payload["ordered_targets"]:
                if target["role"] == "body":
                    values.append("Facts support the first grounded point in this section\nFacts support the second grounded point in this section\nFacts support the third grounded point in this section")
                else:
                    values.append("Grounded title")
            return SimpleNamespace(text=json.dumps({"contents": values}))
        slides = []
        for slide in payload["template_structure"]:
            targets = []
            for target in slide["targets"]:
                content = (f"Facts support the first grounded point in section {slide['slide_number']}\nFacts support the second grounded point in section {slide['slide_number']}\nFacts support the third grounded point in section {slide['slide_number']}" if target["role"] == "body" else f"Grounded title {slide['slide_number']}")
                targets.append({"target_kind": target["target_kind"], "target_id": target["target_id"], "content": content})
            slides.append({"slide_number": slide["slide_number"], "layout_name": slide["layout_name"], "section_type": slide.get("section_type"), "targets": targets})
        return SimpleNamespace(text=json.dumps({"slides": slides}))
    watsonx = MagicMock(); watsonx.chat.side_effect = respond
    embedder = MagicMock(); embedder.embed.return_value = EmbeddingResult(indexed_template["template_embedding"], indexed_template["embedding_model_id"])
    retriever = AutomaticRetriever(indexed_template)
    result = GenerationPipeline(retriever, watsonx, tmp_path, tmp_path / "generated").run_automatic(model_id="generation-model", embedding_client=embedder, topic="Pitch", complete_demand="Build an investor pitch", source_content="Facts", audience="investors", tone="professional", desired_slide_count=1, required_sections=["problem"], visual_preferences="minimal", confidence_threshold=.6)
    assert result.output_path.exists() and retriever.binary_requested
    assert result.stages[-1] == "Temporary files cleaned"
    assert any("below" in warning for warning in result.warnings)
    assert result.generation_model_id == "generation-model"


def test_complete_manual_pipeline_uses_selected_template_without_retrieval(indexed_template, valid_payload, tmp_path):
    class ManualRetriever:
        def __init__(self): self.requested = None
        def get_full(self, template_id): self.requested = template_id; return indexed_template

    retriever = ManualRetriever()
    watsonx = MagicMock()
    def respond(model_id, messages, **kwargs):
        payload = json.loads(messages[1]["content"].split("\n", 1)[1])
        slide = payload["planned_slide"]
        return SimpleNamespace(text=json.dumps({"slides": [{
            "slide_number": slide["slide_number"], "section_type": slide["section_type"],
            "title": slide.get("section_title") or "Grounded title",
            "bullets": ([] if slide["section_type"] in {"title", "qa"} else [
                "Facts establish the central context", "Facts clarify the relevant evidence",
                "Facts support the practical conclusion",
            ]), "speaker_notes": None,
        }]}))
    watsonx.chat.side_effect = respond
    result = GenerationPipeline(retriever, watsonx, tmp_path, tmp_path / "generated").run_manual(
        template_id=indexed_template["template_id"], model_id="generation-model",
        topic="Manual deck", complete_demand="Use selected template", source_content="Facts",
        audience="leaders", tone="professional", desired_slide_count=14,
        required_sections=[], visual_preferences="", additional_instructions="",
    )
    assert retriever.requested == indexed_template["template_id"]
    assert result.output_path.exists() and result.selection is None
    generated = Presentation(result.output_path)
    assert len(generated.slides) == 14
    assert all(
        shape.left >= 0 and shape.top >= 0
        and shape.left + shape.width <= generated.slide_width
        and shape.top + shape.height <= generated.slide_height
        for slide in generated.slides for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    all_text = "\n".join(shape.text for slide in generated.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert "Original title" not in all_text and "Original subtitle" not in all_text
