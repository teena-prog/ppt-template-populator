from __future__ import annotations

import io
from types import SimpleNamespace
import pytest

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from src.document_extractor import extract_text
from src.ppt_populator import populate_presentation
from src.response_models import PresentationContent, SlideContent, TargetContent
from src.response_validator import ResponseValidationError, validate_response
from src.slide_planner import FULL_SLIDE_BLUEPRINT, SLIDE_BLUEPRINT, build_slide_plan, build_ten_slide_plan
from src.style_profile import build_style_profile
from src.template_parser import parse_template
from src.target_metadata import required_targets_by_slide
from src.visual_validator import validate_presentation_layout


def test_slide_plan_has_exactly_ten_roles_and_uses_all_source(template_metadata):
    source = "\n\n".join(f"Fact section {index}. Supporting detail." for index in range(1, 17))
    planned = build_ten_slide_plan(template_metadata, source)
    assert len(planned["slides"]) == 10
    assert [item["role"] for item in planned["slide_plan"]] == [item[0] for item in SLIDE_BLUEPRINT]
    combined = " ".join(item["source_excerpt"] for item in planned["slide_plan"])
    assert all(f"Fact section {index}." in combined for index in range(1, 17))


def test_introduction_receives_multiple_grounded_source_units(template_metadata):
    source = " ".join(f"Grounded statement {index}." for index in range(1, 8))
    planned = build_slide_plan(template_metadata, source, 14)
    introduction = planned["slide_plan"][1]
    assert introduction["role"] == "introduction"
    assert sum(f"Grounded statement {index}." in introduction["source_excerpt"] for index in range(1, 8)) == 5


def test_default_reference_outline_has_fourteen_single_purpose_slides(template_metadata):
    planned = build_slide_plan(template_metadata, "Grounded evidence supports the requested presentation.", 14)
    assert len(planned["slides"]) == 14
    assert [slide["section_type"] for slide in planned["slides"]] == [item[0] for item in FULL_SLIDE_BLUEPRINT]
    assert len(required_targets_by_slide(planned)[1]) == 1
    assert len(required_targets_by_slide(planned)[14]) == 1


def test_mandatory_slots_use_cover_intro_and_closing_layouts(template_metadata):
    planned = build_slide_plan(template_metadata, "Grounded source evidence.", 14)
    assert planned["slide_plan"][0]["source_slide_number"] == 1
    assert planned["slides"][0]["section_type"] == "title"
    assert planned["slides"][1]["section_type"] == "introduction"
    assert planned["slide_plan"][-1]["source_slide_number"] == template_metadata["slides"][-1]["slide_number"]
    assert planned["slides"][-1]["section_type"] == "closing"


def test_generic_style_profile_scales_across_template_dimensions(template_metadata):
    standard = dict(template_metadata, slide_width=12_192_000, slide_height=6_858_000)
    classic = dict(template_metadata, slide_width=9_144_000, slide_height=6_858_000)
    widescreen = build_style_profile(standard)
    four_three = build_style_profile(classic)
    assert round(widescreen.aspect_ratio, 2) == 1.78
    assert round(four_three.aspect_ratio, 2) == 1.33
    assert .035 <= widescreen.horizontal_margin_ratio <= .08
    assert widescreen.title_size_pt > widescreen.body_size_pt >= 16


def test_planned_section_order_separates_problem_objectives_and_solution(template_metadata):
    planned = build_ten_slide_plan(template_metadata, "Grounded source facts support each planned section.")
    assert [slide["section_type"] for slide in planned["slides"]] == [
        "title", "introduction", "problem_statement", "objectives",
        "proposed_solution", "methodology_workflow", "case_study",
        "novelty_key_findings", "conclusion", "closing",
    ]
    assert len(set(slide["section_type"] for slide in planned["slides"])) == 10


def _target(number, target_id, role, kind="shape"):
    return {"slide_number": number, "target_kind": kind, "target_id": target_id,
            "role": role, "existing_text": f"{role.title()} {number}-{target_id}", "approximate_max_characters": 300,
            "max_content_length": 300, "required": True, "replaceable": True,
            "font_size": 32 if role == "title" else 20,
            "position": {"left": 10, "top": 10 if role == "title" else 100, "width": 500, "height": 60 if role == "title" else 300}}


def test_title_and_qa_accept_title_only_and_content_gets_safe_body_fallback():
    template = {"template_id": "title-only", "slides": [{"slide_number": 1, "layout_name": "Blank", "shapes": [{"shape_id": 7}], "targets": [_target(1, 7, "title")]}]}
    planned = build_ten_slide_plan(template, "Grounded source content supports the presentation.")
    required = required_targets_by_slide(planned)
    assert len(required[1]) == 1
    assert len(required[10]) == 1
    for number in range(2, 10):
        roles = {metadata["role"] for metadata in required[number].values()}
        assert roles == {"title", "body"}
        assert planned["slide_plan"][number - 1]["create_body_target"] is not None


def test_optional_subtitle_is_not_required_on_title_slide():
    template = {"template_id": "subtitle", "slides": [{"slide_number": 1, "layout_name": "Title", "shapes": [{"shape_id": 2}, {"shape_id": 3}], "targets": [_target(1, 2, "title"), _target(1, 3, "subtitle")]}]}
    planned = build_ten_slide_plan(template, "Grounded source content supports the presentation.")
    assert len(required_targets_by_slide(planned)[1]) == 1


def test_content_slides_prefer_existing_compatible_canva_layout():
    template = {"template_id": "canva", "slides": [
        {"slide_number": 1, "layout_name": "Blank", "shapes": [{"shape_id": 2}], "targets": [_target(1, 2, "title")]},
        {"slide_number": 2, "layout_name": "Blank", "shapes": [{"shape_id": 8}, {"shape_id": 9}], "targets": [_target(2, 8, "title"), _target(2, 9, "body")]},
    ]}
    planned = build_ten_slide_plan(template, "Grounded source content supports the presentation.")
    assert all(slide["source_slide_number"] == 2 for slide in planned["slides"][1:9])
    assert all(item["create_body_target"] is None for item in planned["slide_plan"][1:9])
    assert all({target["target_kind"] for target in slide["targets"] if target["required"]} == {"shape"} for slide in planned["slides"][1:9])


def test_planned_body_requires_sufficient_bullets_and_rejects_mixed_purpose(template_metadata):
    planned = build_ten_slide_plan(template_metadata, "Teams use evidence to define objectives and improve decisions.")
    required = required_targets_by_slide(planned)
    slides = []
    for slide in planned["slides"]:
        targets = []
        for key, metadata in required[slide["slide_number"]].items():
            metadata["approximate_max_characters"] = metadata["max_content_length"] = 300
            if metadata["role"] == "body":
                number = slide["slide_number"]
                content = f"Evidence guides teams toward clearer decisions for section {number}\nTeams use evidence for focused planning in assigned section {number}\nEvidence helps teams explain decisions for this distinct section {number}"
            else:
                content = f"Section {slide['slide_number']}"
            targets.append({"target_kind": key[0], "target_id": key[1], "content": content})
        slides.append({"slide_number": slide["slide_number"], "layout_name": slide["layout_name"], "section_type": slide["section_type"], "targets": targets})
    validate_response({"slides": slides}, planned, presentation_title="Grounded deck")
    body = next(target for target in slides[2]["targets"] if required[3][(target["target_kind"], target["target_id"])]["role"] == "body")
    body["content"] = "Objectives: combine another purpose here\nOnly two bullets are present on this slide"
    _, warnings = validate_response({"slides": slides}, planned, presentation_title="Grounded deck")
    assert any("prefers 3-5 bullets" in warning for warning in warnings)


def test_pdf_extraction_removes_repeated_headers_and_footers_without_truncation(monkeypatch):
    pages = [
        SimpleNamespace(extract_text=lambda n=n: f"Annual Report\nPage heading {n}\nUnique fact {n}.\nConfidential" )
        for n in range(1, 5)
    ]
    monkeypatch.setattr("pypdf.PdfReader", lambda stream: SimpleNamespace(pages=pages, metadata=None))
    result = extract_text(b"fake-pdf", "report.pdf")
    assert "Annual Report" not in result.text
    assert "Confidential" not in result.text
    assert all(f"Unique fact {n}." in result.text for n in range(1, 5))
    assert not any("truncated" in warning for warning in result.warnings)


def test_population_builds_ten_slides_replaces_text_and_preserves_logo(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Sample Company"
    slide.placeholders[1].text = "Lorem ipsum sample text"
    image = Image.new("RGB", (20, 20), "red")
    blob = io.BytesIO(); image.save(blob, format="PNG"); blob.seek(0)
    logo = slide.shapes.add_picture(blob, Inches(0.1), Inches(0.1), Inches(0.25), Inches(0.25))
    logo.name = "Brand Logo"
    stock_image = Image.new("RGB", (80, 60), "blue")
    stock_blob = io.BytesIO(); stock_image.save(stock_blob, format="PNG"); stock_blob.seek(0)
    stock = slide.shapes.add_picture(stock_blob, Inches(8), Inches(2), Inches(1.5), Inches(1.2))
    stock.name = "Stock Photo"
    source = tmp_path / "source.pptx"; prs.save(source)
    metadata = parse_template(source, file_path="templates/source.pptx")
    planned = build_ten_slide_plan(metadata, "Grounded fact one. Grounded fact two.")
    required = required_targets_by_slide(planned)
    slides = []
    for item in planned["slides"]:
        targets = [TargetContent(target_kind=key[0], target_id=key[1], content=f"New {target['role']}") for key, target in required[item["slide_number"]].items()]
        slides.append(SlideContent(slide_number=item["slide_number"], layout_name=item["layout_name"], targets=targets))
    output, _ = populate_presentation(source, PresentationContent(presentation_title="Deck", slides=slides), tmp_path / "out", planned)
    generated = Presentation(output)
    assert len(generated.slides) == 10
    assert all("Sample Company" not in " ".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)) for slide in generated.slides)
    assert all(any(getattr(shape, "name", "") == "Brand Logo" for shape in slide.shapes) for slide in generated.slides)
    assert all(not any(getattr(shape, "name", "") == "Stock Photo" for shape in slide.shapes) for slide in generated.slides)


def test_visual_validation_ignores_static_small_print_and_planned_title(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(8), Inches(.6))
    title.text = "Problem Statement"
    footer = slide.shapes.add_textbox(Inches(.5), Inches(6.8), Inches(3), Inches(.2))
    footer.text = "BRAND NAME"
    footer.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    path = tmp_path / "visual.pptx"; prs.save(path)
    template = {"slide_count": 1, "slides": [{"slide_number": 1, "layout_name": "Blank",
        "planned_section_title": "Problem Statement", "targets": [
            {"target_kind": "shape", "target_id": title.shape_id, "role": "title", "existing_text": "Problem Statement",
             "replaceable": True, "required": True, "approximate_max_characters": 80},
            {"target_kind": "shape", "target_id": footer.shape_id, "role": "footer", "existing_text": "BRAND NAME",
             "replaceable": False, "required": False, "approximate_max_characters": 40},
        ]}]}
    categories = {issue.category for issue in validate_presentation_layout(path, template)}
    assert "sample_text" not in categories
    assert "font_size" not in categories
